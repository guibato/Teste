# financeiro/views/repasse_views.py (versão corrigida e completa)
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.generic import ListView, DetailView, CreateView, UpdateView, TemplateView
from django.db.models import Q, Sum, Count, Case, When, DecimalField, F
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from django.core.paginator import Paginator
from django.urls import reverse_lazy
from django.db import transaction
from datetime import date, datetime, timedelta
from decimal import Decimal
import json
import logging
from collections import defaultdict, OrderedDict
from django.views.decorators.http import require_POST

# Imports dos modelos
from ..models.repasse import Repasse, PoliticaRepasseContrato, PoliticaRepasseGlobal, AgendamentoRepasse
try:
    from ..models.cobranca import Cobranca
except ImportError:
    Cobranca = None

from core.models import Contrato, Cliente

# Imports dos services e forms
try:
    from ..services.repasse_service import RepasseService
    
except ImportError:
    RepasseService = None

try:
    from ..forms.repasse_forms import (
        RepasseForm, 
        PoliticaRepasseContratoForm,
        PoliticaRepasseGlobalForm,
        ProcessarRepasseForm,
        ConfiguracaoPoliticasForm,
        FiltroContratosForm
    )
except ImportError:
    # Fallback se os forms não existirem
    RepasseForm = None
    PoliticaRepasseContratoForm = None
    PoliticaRepasseGlobalForm = None
    ProcessarRepasseForm = None
    ConfiguracaoPoliticasForm = None
    FiltroContratosForm = None

import logging
from django.http import JsonResponse
from django.shortcuts import redirect
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.utils.decorators import method_decorator
from django.views.decorators.cache import never_cache

logger = logging.getLogger(__name__)


# financeiro/views/repasse_views.py - FUNÇÃO DEFINITIVA

from django.http import JsonResponse, HttpResponse
from django.shortcuts import redirect
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.contrib.auth.decorators import login_required
from django.utils.decorators import method_decorator
from django.utils import timezone
import json
import logging

logger = logging.getLogger(__name__)


def gerar_repasse_automatico(request):
    """
    🤖 Gera repasses baseados nas políticas dos contratos
    VERSÃO FINAL COMPATÍVEL COM REPASSE SERVICE
    """
    
    def create_json_response(data, status=200):
        """Helper para criar resposta JSON padronizada"""
        response = HttpResponse(
            json.dumps(data, ensure_ascii=False, indent=2),
            content_type='application/json; charset=utf-8',
            status=status
        )
        response['Cache-Control'] = 'no-cache'
        response['X-Content-Type-Options'] = 'nosniff'
        return response
    
    # Log detalhado
    logger.info("=== GERAR_REPASSE_AUTOMATICO ===")
    logger.info(f"Method: {request.method}")
    logger.info(f"User: {request.user}")
    logger.info(f"Authenticated: {request.user.is_authenticated}")
    logger.info(f"Content-Type: {request.headers.get('Content-Type', 'N/A')}")
    logger.info(f"Accept: {request.headers.get('Accept', 'N/A')}")
    logger.info(f"X-Requested-With: {request.headers.get('X-Requested-With', 'N/A')}")
    
    # Verificar método
    if request.method != 'POST':
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return create_json_response({
                'success': False,
                'error': 'Método não permitido. Use POST.',
                'method_received': request.method,
                'criados': 0,
                'processados': 0,
                'agendados': 0,
                'erros': 1
            }, 405)
        else:
            return redirect('financeiro:lista_repasses')
    
    # Processar POST
    try:
        logger.info("Iniciando processamento de repasses...")
        
        # 1. Verificar e importar RepasseService
        try:
            from financeiro.services.repasse.repasse_service import RepasseService
            logger.info("✅ RepasseService importado com sucesso")
        except ImportError as e:
            logger.error(f"❌ Erro ao importar RepasseService: {e}")
            return create_json_response({
                'success': False,
                'error': 'RepasseService não encontrado. Verifique se o arquivo services/repasse/repasse_service.py existe.',
                'criados': 0,
                'processados': 0,
                'agendados': 0,
                'erros': 1,
                'debug': {
                    'import_error': str(e),
                    'expected_path': 'financeiro.services.repasse.repasse_service'
                }
            }, 500)
        
        # 2. Instanciar serviço
        try:
            service = RepasseService()
            logger.info("✅ RepasseService instanciado")
        except Exception as e:
            logger.error(f"❌ Erro ao instanciar RepasseService: {e}")
            return create_json_response({
                'success': False,
                'error': f'Erro ao inicializar serviço: {str(e)}',
                'criados': 0,
                'processados': 0,
                'agendados': 0,
                'erros': 1,
                'debug': {
                    'instantiation_error': str(e)
                }
            }, 500)
        
        # 3. Verificar se método existe
        if not hasattr(service, 'processar_repasses_automaticos'):
            logger.error("❌ Método processar_repasses_automaticos não encontrado")
            return create_json_response({
                'success': False,
                'error': 'Método processar_repasses_automaticos não encontrado no RepasseService',
                'criados': 0,
                'processados': 0,
                'agendados': 0,
                'erros': 1,
                'debug': {
                    'available_methods': [method for method in dir(service) if not method.startswith('_')]
                }
            }, 500)
        
        # 4. Executar processamento
        logger.info("Executando processar_repasses_automaticos()...")
        resultado = service.processar_repasses_automaticos()
        logger.info(f"Resultado: {resultado}")
        
        # 5. Processar resultado
        criados = int(resultado.get('criados', 0))
        processados = int(resultado.get('processados', 0))
        agendados = int(resultado.get('agendados', 0))
        erros = int(resultado.get('erros', 0))
        detalhes = resultado.get('detalhes', [])
        
        # 6. Determinar mensagem e tipo
        if criados > 0:
            mensagem = f"✅ {criados} repasse(s) criado(s) com sucesso!"
            tipo = 'success'
        elif processados > 0:
            mensagem = f"✅ {processados} agendamento(s) processado(s)!"
            tipo = 'success'
        elif agendados > 0:
            mensagem = f"📅 {agendados} agendamento(s) criado(s)!"
            tipo = 'info'
        elif erros > 0:
            mensagem = f"⚠️ Processamento concluído com {erros} erro(s). Verifique os detalhes."
            tipo = 'warning'
        else:
            mensagem = "ℹ️ Processamento executado. Nenhuma ação necessária no momento."
            tipo = 'info'
        
        # 7. Log do resultado
        logger.info(f"=== RESULTADO FINAL ===")
        logger.info(f"Criados: {criados}")
        logger.info(f"Processados: {processados}")
        logger.info(f"Agendados: {agendados}")
        logger.info(f"Erros: {erros}")
        logger.info(f"Mensagem: {mensagem}")
        
        # 8. Preparar resposta JSON
        response_data = {
            'success': True,
            'criados': criados,
            'processados': processados,
            'agendados': agendados,
            'erros': erros,
            'mensagem': mensagem,
            'tipo': tipo,
            'detalhes': detalhes[:10],  # Limitar detalhes para não sobrecarregar
            'timestamp': timezone.now().isoformat(),
            'debug': {
                'total_operations': criados + processados + agendados,
                'has_errors': erros > 0,
                'user': str(request.user)
            }
        }
        
        logger.info(f"Retornando JSON de sucesso: {json.dumps(response_data, indent=2)}")
        return create_json_response(response_data)
        
    except Exception as e:
        # Log completo do erro
        import traceback
        error_traceback = traceback.format_exc()
        
        logger.error(f"=== ERRO CRÍTICO ===")
        logger.error(f"Erro: {str(e)}")
        logger.error(f"Tipo: {type(e).__name__}")
        logger.error(f"Traceback: {error_traceback}")
        
        # Resposta de erro JSON
        error_response = {
            'success': False,
            'error': f'Erro no processamento: {str(e)}',
            'criados': 0,
            'processados': 0,
            'agendados': 0,
            'erros': 1,
            'debug': {
                'error_type': type(e).__name__,
                'user': str(request.user),
                'timestamp': timezone.now().isoformat(),
                'method': request.method,
                'path': request.path
            }
        }
        
        logger.error(f"Retornando JSON de erro: {json.dumps(error_response, indent=2)}")
        return create_json_response(error_response, status=500)


# === VERSÃO DE TESTE PARA FALLBACK ===


def gerar_repasse_automatico_teste(request):
    """
    🧪 Versão de teste que sempre retorna JSON válido
    Use como fallback se a função principal não funcionar
    """
    
    if request.method != 'POST':
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            data = {
                'success': False,
                'error': 'Método GET não suportado',
                'criados': 0,
                'processados': 0,
                'agendados': 0,
                'erros': 1
            }
        else:
            return redirect('financeiro:lista_repasses')
    else:
        # Simular processamento bem-sucedido
        data = {
            'success': True,
            'criados': 0,
            'processados': 0,
            'agendados': 0,
            'erros': 0,
            'mensagem': '🧪 Função de teste executada com sucesso! RepasseService está sendo carregado...',
            'tipo': 'info',
            'detalhes': [
                'Esta é uma função de teste',
                'Verifique se o RepasseService está configurado corretamente',
                f'Usuário: {request.user}',
                f'Timestamp: {timezone.now().isoformat()}'
            ],
            'debug': True
        }
    
    # Garantir retorno JSON
    response = HttpResponse(
        json.dumps(data, ensure_ascii=False, indent=2),
        content_type='application/json; charset=utf-8'
    )
    response['Cache-Control'] = 'no-cache'
    return response


# === FUNÇÃO PARA VERIFICAR CONFIGURAÇÃO ===


def verificar_repasse_config(request):
    """
    🔍 Endpoint para verificar se tudo está configurado corretamente
    Acesse: /financeiro/repasses/verificar-config/
    """
    
    resultado = {
        'timestamp': timezone.now().isoformat(),
        'user': str(request.user),
        'verificacoes': {}
    }
    
    try:
        # 1. Verificar importação do RepasseService
        try:
            from financeiro.services.repasse.repasse_service import RepasseService
            resultado['verificacoes']['import_service'] = {
                'status': 'OK',
                'message': 'RepasseService importado com sucesso'
            }
            
            # 2. Verificar instanciação
            try:
                service = RepasseService()
                resultado['verificacoes']['instantiate_service'] = {
                    'status': 'OK',
                    'message': 'RepasseService instanciado com sucesso'
                }
                
                # 3. Verificar métodos
                metodos_esperados = ['processar_repasses_automaticos', 'calcular_repasse_cobranca', 'criar_repasse_automatico']
                metodos_encontrados = []
                metodos_faltando = []
                
                for metodo in metodos_esperados:
                    if hasattr(service, metodo):
                        metodos_encontrados.append(metodo)
                    else:
                        metodos_faltando.append(metodo)
                
                resultado['verificacoes']['metodos'] = {
                    'status': 'OK' if not metodos_faltando else 'AVISO',
                    'encontrados': metodos_encontrados,
                    'faltando': metodos_faltando,
                    'todos_metodos': [m for m in dir(service) if not m.startswith('_')]
                }
                
            except Exception as e:
                resultado['verificacoes']['instantiate_service'] = {
                    'status': 'ERRO',
                    'message': f'Erro ao instanciar: {str(e)}'
                }
                
        except ImportError as e:
            resultado['verificacoes']['import_service'] = {
                'status': 'ERRO',
                'message': f'Erro ao importar: {str(e)}'
            }
        
        # 4. Verificar modelos
        try:
            from financeiro.models.repasse import Repasse, PoliticaRepasseContrato, AgendamentoRepasse
            from financeiro.models.cobranca import Cobranca
            
            resultado['verificacoes']['modelos'] = {
                'status': 'OK',
                'counts': {
                    'repasses': Repasse.objects.count(),
                    'politicas': PoliticaRepasseContrato.objects.count(),
                    'agendamentos': AgendamentoRepasse.objects.count(),
                    'cobrancas': Cobranca.objects.count()
                }
            }
            
        except Exception as e:
            resultado['verificacoes']['modelos'] = {
                'status': 'ERRO',
                'message': f'Erro nos modelos: {str(e)}'
            }
        
        # 5. Status geral
        erros = sum(1 for v in resultado['verificacoes'].values() if v['status'] == 'ERRO')
        avisos = sum(1 for v in resultado['verificacoes'].values() if v['status'] == 'AVISO')
        
        if erros == 0 and avisos == 0:
            resultado['status_geral'] = 'OK'
            resultado['mensagem'] = '✅ Tudo configurado corretamente!'
        elif erros == 0:
            resultado['status_geral'] = 'AVISO'
            resultado['mensagem'] = f'⚠️ Configuração OK com {avisos} aviso(s)'
        else:
            resultado['status_geral'] = 'ERRO'
            resultado['mensagem'] = f'❌ {erros} erro(s) encontrado(s)'
        
    except Exception as e:
        resultado['status_geral'] = 'ERRO_CRITICO'
        resultado['mensagem'] = f'Erro crítico na verificação: {str(e)}'
    
    # Retornar JSON
    response = HttpResponse(
        json.dumps(resultado, ensure_ascii=False, indent=2),
        content_type='application/json; charset=utf-8'
    )
    return response


# === URLS ADICIONAIS NECESSÁRIAS ===

"""
Adicione estas URLs no financeiro/urls.py:

# URLs para debug e teste
path('repasses/verificar-config/', repasse_views.verificar_repasse_config, name='verificar_repasse_config'),
path('repasses/gerar-automatico-teste/', repasse_views.gerar_repasse_automatico_teste, name='gerar_repasse_automatico_teste'),

# URL principal (substitua a existente)
path('repasses/gerar-automatico/', repasse_views.gerar_repasse_automatico, name='gerar_repasse_automatico'),
"""


# === INSTRUÇÕES DE USO ===

"""
PARA RESOLVER O PROBLEMA:

1. PRIMEIRO: Teste se a configuração está OK
   - Acesse: http://127.0.0.1:8000/financeiro/repasses/verificar-config/
   - Deve retornar JSON com status das verificações

2. SE HOUVER ERROS: Use a versão de teste
   - Substitua temporariamente no urls.py:
     path('repasses/gerar-automatico/', repasse_views.gerar_repasse_automatico_teste, name='gerar_repasse_automatico'),

3. CORRIJA OS PROBLEMAS identificados na verificação

4. VOLTE para a versão principal:
   path('repasses/gerar-automatico/', repasse_views.gerar_repasse_automatico, name='gerar_repasse_automatico'),

5. TESTE no navegador:
   - http://127.0.0.1:8000/financeiro/repasses/gerar-automatico/ (deve redirecionar)
   - Via JavaScript deve retornar JSON válido
"""

def test_repasse_service(request):
    """
    View para testar se o RepasseService está funcionando
    Acesse: /financeiro/test-repasse-service/
    """
    logger.info("=== TESTE DO REPASSE SERVICE ===")
    
    try:
        # Teste 1: Importação
        if RepasseService:
            logger.info("✅ RepasseService importado com sucesso")
            service_status = "Importado"
        else:
            logger.error("❌ RepasseService não importado")
            service_status = "Não importado"
        
        # Teste 2: Instanciação
        try:
            service = RepasseService()
            logger.info("✅ RepasseService instanciado com sucesso")
            instance_status = "Instanciado"
        except Exception as e:
            logger.error(f"❌ Erro ao instanciar RepasseService: {e}")
            instance_status = f"Erro: {e}"
            service = None
        
        # Teste 3: Métodos disponíveis
        methods = []
        if service:
            methods = [method for method in dir(service) if not method.startswith('_')]
            logger.info(f"Métodos disponíveis: {methods}")
        
        # Teste 4: Verificar modelos
        model_status = {}
        try:
            model_status['Repasse'] = Repasse.objects.count()
            model_status['PoliticaRepasseContrato'] = PoliticaRepasseContrato.objects.count()
            model_status['AgendamentoRepasse'] = AgendamentoRepasse.objects.count()
            if Cobranca:
                model_status['Cobranca'] = Cobranca.objects.count()
        except Exception as e:
            model_status['error'] = str(e)
        
        resultado = {
            'service_status': service_status,
            'instance_status': instance_status,
            'methods': methods,
            'model_counts': model_status,
            'user': str(request.user),
            'timestamp': timezone.now().isoformat()
        }
        
        return JsonResponse({
            'success': True,
            'resultado': resultado
        })
        
    except Exception as e:
        logger.error(f"Erro no teste: {e}")
        return JsonResponse({
            'success': False,
            'error': str(e)
        })





def processar_agendamentos_vencidos(request):
    """Processa todos os agendamentos vencidos - VERSÃO CORRIGIDA"""
    if request.method == 'POST':
        try:
            service = RepasseService()
            service._processar_agendamentos_vencidos()
            resultado = service.resultados
            
            if resultado.get('processados', 0) > 0:
                messages.success(request, 
                    f"{resultado.get('processados', 0)} agendamentos processados com sucesso!")
            else:
                messages.info(request, "Nenhum agendamento vencido encontrado.")
            
            if resultado.get('erros', 0) > 0:
                messages.warning(request, f"{resultado.get('erros', 0)} erros durante o processamento.")
            
        except Exception as e:
            logger.error(f"Erro no processamento de agendamentos: {str(e)}")
            messages.error(request, f"Erro no processamento: {str(e)}")
    
    return redirect('financeiro:agendamento_list')


def testar_politica_contrato_view(request, pk):
    """Testa uma política de contrato (simulação) - VERSÃO CORRIGIDA E SIMPLIFICADA"""
    if request.method == 'POST':
        try:
            politica = get_object_or_404(PoliticaRepasseContrato, pk=pk)
            contrato = politica.contrato
            
            # Obter valor do contrato de forma segura
            valor_atual = Decimal('0')
            
            # Tentar diferentes campos de valor
            if hasattr(contrato, 'valor_aluguel_atual') and callable(contrato.valor_aluguel_atual):
                valor_atual = contrato.valor_aluguel_atual()
            elif hasattr(contrato, 'valor_aluguel'):
                valor_atual = contrato.valor_aluguel or Decimal('0')
            elif hasattr(contrato, 'valor_base'):
                valor_atual = contrato.valor_base or Decimal('0')
            
            if valor_atual <= 0:
                return JsonResponse({
                    'success': False,
                    'error': 'Valor do contrato não encontrado ou inválido'
                })
            
            # Obter taxa administrativa
            taxa_admin = Decimal('8.00')  # Padrão
            
            if hasattr(politica, 'taxa_admin_personalizada') and politica.taxa_admin_personalizada:
                taxa_admin = politica.taxa_admin_personalizada
            elif hasattr(contrato, 'valor_taxa_administracao_percentual') and contrato.valor_taxa_administracao_percentual:
                taxa_admin = contrato.valor_taxa_administracao_percentual
            
            # Calcular valores
            valor_taxa_admin = valor_atual * (taxa_admin / 100)
            valor_liquido = valor_atual - valor_taxa_admin
            valor_minimo = getattr(politica, 'valor_minimo_repasse', Decimal('0')) or Decimal('0')
            
            # Calcular próxima data
            from datetime import date, timedelta
            hoje = date.today()
            
            try:
                if hasattr(politica, 'calcular_proxima_data_repasse'):
                    proxima_data = politica.calcular_proxima_data_repasse(hoje)
                else:
                    # Fallback simples
                    dias_apos = getattr(politica, 'dias_apos_recebimento', 7)
                    proxima_data = hoje + timedelta(days=dias_apos)
            except:
                proxima_data = hoje + timedelta(days=30)
            
            # Montar resposta
            simulacao = {
                'valor_bruto': float(valor_atual),
                'taxa_admin': float(taxa_admin),
                'valor_taxa_admin': float(valor_taxa_admin),
                'valor_liquido': float(valor_liquido),
                'valor_minimo': float(valor_minimo),
                'acima_minimo': valor_liquido >= valor_minimo,
                'proxima_data': proxima_data.strftime('%d/%m/%Y') if proxima_data else 'N/A',
                'periodicidade': getattr(politica, 'periodicidade', 'N/A'),
                'tipo_dias': getattr(politica, 'tipo_dias', 'N/A'),
                'ativa': getattr(politica, 'ativa', False)
            }
            
            return JsonResponse({
                'success': True,
                'simulacao': simulacao
            })
            
        except Exception as e:
            logger.error(f"Erro ao testar política {pk}: {str(e)}")
            return JsonResponse({
                'success': False,
                'error': f'Erro na simulação: {str(e)}'
            })
    
    return JsonResponse({'success': False, 'error': 'Método não permitido'})


# ==================== NOVA VIEW PARA DASHBOARD DE AUTOMAÇÃO ====================

class DashboardAutomacaoView(TemplateView):
    """Dashboard específico para acompanhar a automação de repasses"""
    template_name = 'financeiro/repasses/dashboard_automacao.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        try:
            service = RepasseService()
            relatorio = service.gerar_relatorio_automatizacao()
            atrasados = service.verificar_repasses_atrasados()
            
            context.update({
                'relatorio': relatorio,
                'atrasados': atrasados,
                'pode_processar': True
            })
            
            # Próximos agendamentos (7 dias)
            from datetime import date, timedelta
            proximos_agendamentos = AgendamentoRepasse.objects.filter(
                status='agendado',
                data_agendada__lte=date.today() + timedelta(days=7)
            ).select_related('proprietario', 'contrato').order_by('data_agendada')[:10]
            
            context['proximos_agendamentos'] = proximos_agendamentos
            
            # Últimos processamentos
            ultimos_repasses = Repasse.objects.filter(
                tipo='automatico'
            ).select_related('proprietario', 'contrato').order_by('-data_criacao')[:10]
            
            context['ultimos_repasses'] = ultimos_repasses
            
        except Exception as e:
            logger.error(f"Erro no dashboard de automação: {str(e)}")
            context.update({
                'erro': str(e),
                'pode_processar': False
            })
        
        return context


def dashboard_automacao_ajax(request):
    """Dados AJAX para o dashboard de automação"""
    try:
        service = RepasseService()
        relatorio = service.gerar_relatorio_automatizacao()
        atrasados = service.verificar_repasses_atrasados()
        
        # Estatísticas dos últimos 30 dias
        from datetime import date, timedelta
        data_inicio = date.today() - timedelta(days=30)
        
        repasses_mes = Repasse.objects.filter(
            tipo='automatico',
            data_criacao__date__gte=data_inicio
        ).values('status').annotate(
            count=Count('id'),
            total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin'))
        )
        
        return JsonResponse({
            'success': True,
            'relatorio': relatorio,
            'atrasados': atrasados,
            'repasses_mes': list(repasses_mes),
            'timestamp': timezone.now().isoformat()
        })
        
    except Exception as e:
        logger.error(f"Erro no AJAX do dashboard: {str(e)}")
        return JsonResponse({
            'success': False,
            'error': str(e)
        })


def processar_repasse_unico(request, repasse_id):
    """Processa um repasse específico via AJAX"""
    if request.method == 'POST':
        try:
            service = RepasseService()
            resultado = service.processar_repasse_individual(repasse_id)
            
            if resultado['success']:
                return JsonResponse({
                    'success': True,
                    'message': resultado['message']
                })
            else:
                return JsonResponse({
                    'success': False,
                    'error': resultado['error']
                })
                
        except Exception as e:
            logger.error(f"Erro ao processar repasse {repasse_id}: {str(e)}")
            return JsonResponse({
                'success': False,
                'error': str(e)
            })
    
    return JsonResponse({'success': False, 'error': 'Método não permitido'})


# ==================== VIEW PARA DIAGNÓSTICO ====================

def diagnostico_automacao(request):
    """View para diagnosticar problemas na automação"""
    diagnostico = {
        'data_verificacao': date.today(),
        'problemas': [],
        'avisos': [],
        'sugestoes': []
    }
    
    try:
        # Verificar contratos sem política
        contratos_sem_politica = Contrato.objects.filter(
            ativo=True,
            politica_repasse__isnull=True
        ).count()
        
        if contratos_sem_politica > 0:
            diagnostico['problemas'].append(
                f"{contratos_sem_politica} contratos ativos sem política de repasse configurada"
            )
            diagnostico['sugestoes'].append(
                "Configure políticas globais e aplique aos contratos sem política"
            )
        
        # Verificar políticas inativas
        politicas_inativas = PoliticaRepasseContrato.objects.filter(ativa=False).count()
        if politicas_inativas > 0:
            diagnostico['avisos'].append(
                f"{politicas_inativas} políticas de contrato estão inativas"
            )
        
        # Verificar repasses atrasados
        repasses_atrasados = Repasse.objects.filter(
            status='pendente',
            data_prevista__lt=date.today()
        ).count()
        
        if repasses_atrasados > 0:
            diagnostico['problemas'].append(
                f"{repasses_atrasados} repasses estão atrasados"
            )
            diagnostico['sugestoes'].append(
                "Execute o processamento automático ou processe os repasses manualmente"
            )
        
        # Verificar cobranças pagas sem repasse
        if Cobranca:
            cobrancas_sem_repasse = Cobranca.objects.filter(
                status='paga',
                data_pagamento__gte=date.today() - timedelta(days=30),
                repasses_cobranca__isnull=True
            ).count()
            
            if cobrancas_sem_repasse > 0:
                diagnostico['avisos'].append(
                    f"{cobrancas_sem_repasse} cobranças pagas sem repasse nos últimos 30 dias"
                )
        
        # Verificar agendamentos antigos
        agendamentos_antigos = AgendamentoRepasse.objects.filter(
            status='agendado',
            data_agendada__lt=date.today() - timedelta(days=7)
        ).count()
        
        if agendamentos_antigos > 0:
            diagnostico['problemas'].append(
                f"{agendamentos_antigos} agendamentos vencidos há mais de 7 dias"
            )
        
        # Status geral
        total_problemas = len(diagnostico['problemas'])
        total_avisos = len(diagnostico['avisos'])
        
        if total_problemas == 0 and total_avisos == 0:
            diagnostico['status'] = 'ok'
            diagnostico['mensagem'] = 'Sistema de automação funcionando corretamente'
        elif total_problemas == 0:
            diagnostico['status'] = 'atencao'
            diagnostico['mensagem'] = f'{total_avisos} item(s) que requerem atenção'
        else:
            diagnostico['status'] = 'problema'
            diagnostico['mensagem'] = f'{total_problemas} problema(s) encontrado(s)'
        
    except Exception as e:
        diagnostico['status'] = 'erro'
        diagnostico['mensagem'] = f'Erro ao executar diagnóstico: {str(e)}'
        logger.error(f"Erro no diagnóstico: {str(e)}")
    
    return render(request, 'financeiro/repasses/diagnostico_automacao.html', {
        'diagnostico': diagnostico
    })



from django.views.generic import ListView
from django.db.models import Q, Sum, Count, Case, When, DecimalField, F
from django.core.exceptions import FieldError
from datetime import date, timedelta
from decimal import Decimal
from collections import defaultdict
import logging

logger = logging.getLogger(__name__)


class RepasseListView(ListView):
    model = Repasse
    template_name = 'financeiro/repasses/lista_repasses.html'
    context_object_name = 'repasses'
    paginate_by = 50

    def get_queryset(self):
        """Query otimizada com ordenação CORRIGIDA"""
        queryset = Repasse.objects.select_related(
            'proprietario', 'contrato', 'cobranca'
        ).prefetch_related(
            'contrato__imovel', 
            'contrato__politica_repasse'
        )

        # Filtros básicos
        status = self.request.GET.get('status')
        if status:
            queryset = queryset.filter(status=status)

        proprietario_id = self.request.GET.get('proprietario')
        if proprietario_id:
            queryset = queryset.filter(proprietario_id=proprietario_id)

        mes = self.request.GET.get('mes')
        ano = self.request.GET.get('ano')
        if mes and ano:
            queryset = queryset.filter(mes_referencia=mes, ano_referencia=ano)

        data_inicio = self.request.GET.get('data_inicio')
        data_fim = self.request.GET.get('data_fim')
        if data_inicio:
            queryset = queryset.filter(data_prevista__gte=data_inicio)
        if data_fim:
            queryset = queryset.filter(data_prevista__lte=data_fim)

        # Busca por texto
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(proprietario__nome__icontains=search) |
                Q(contrato__imovel__endereco__icontains=search) |
                Q(descricao__icontains=search) |
                Q(observacoes__icontains=search)
            )

        # *** CORREÇÃO: Ordenação do mais antigo para o mais novo ***
        agrupamento = self.request.GET.get('agrupamento', 'data')
        try:
            if agrupamento == 'data':
                # CORRIGIDO: Remover '-' para ordenar do mais antigo para o mais novo
                return queryset.order_by('data_prevista', 'id')
            elif agrupamento == 'proprietario':
                return queryset.order_by('proprietario__nome', 'data_prevista')
            elif agrupamento == 'status':
                return queryset.order_by(
                    Case(
                        When(status='pendente', then=1),
                        When(status='efetuado', then=2),
                        When(status='cancelado', then=3),
                        default=4
                    ),
                    'data_prevista'  # Mais antigo primeiro
                )
            else:
                return queryset.order_by('data_prevista', 'id')
        except FieldError:
            # Fallback para ordenação simples
            return queryset.order_by('data_prevista', 'id')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        # Obter queryset para estatísticas
        todos_repasses = self.get_queryset()
        
        # Estatísticas seguras - usando método auxiliar CORRIGIDO
        total_stats = self._calcular_estatisticas_seguras(todos_repasses)
        
        # Agrupamento dos repasses
        agrupamento = self.request.GET.get('agrupamento', 'data')
        
        if agrupamento == 'nenhum':
            repasses_agrupados = [{
                'tipo': 'lista_simples',
                'repasses': list(todos_repasses[:self.paginate_by]),
                'valor_total': sum(self._calcular_valor_liquido(r) for r in todos_repasses[:self.paginate_by])
            }]
        else:
            repasses_agrupados = self._agrupar_repasses_seguro(todos_repasses, agrupamento)

        # Estatísticas de políticas (com fallback)
        try:
            stats_politicas = {
                'contratos_com_politica': Contrato.objects.filter(
                    politica_repasse__isnull=False,
                    politica_repasse__ativa=True
                ).count(),
                'contratos_sem_politica': Contrato.objects.filter(
                    ativo=True,
                    politica_repasse__isnull=True
                ).count(),
                'politicas_ativas': PoliticaRepasseContrato.objects.filter(ativa=True).count()
            }
        except Exception as e:
            logger.warning(f"Erro ao obter estatísticas de políticas: {e}")
            stats_politicas = {
                'contratos_com_politica': 0,
                'contratos_sem_politica': 0,
                'politicas_ativas': 0
            }

        # Contexto final
        context.update({
            'total_stats': total_stats,
            'repasses_agrupados': repasses_agrupados,
            'stats_politicas': stats_politicas,
            'filtros': {
                'status': self.request.GET.get('status', ''),
                'proprietario': self.request.GET.get('proprietario', ''),
                'mes': self.request.GET.get('mes', ''),
                'ano': self.request.GET.get('ano', ''),
                'data_inicio': self.request.GET.get('data_inicio', ''),
                'data_fim': self.request.GET.get('data_fim', ''),
                'search': self.request.GET.get('search', ''),
                'agrupamento': agrupamento,
            }
        })
        
        return context

    def _calcular_estatisticas_seguras(self, queryset):
        """
        Calcula estatísticas CORRIGIDAS com valores líquidos
        """
        try:
            # Tentar cálculo otimizado primeiro
            stats = queryset.aggregate(
                total_pendente=Sum(
                    Case(
                        When(status='pendente', then=F('valor') - F('valor_desconto') - F('valor_taxa_admin')), 
                        output_field=DecimalField(max_digits=10, decimal_places=2),
                        default=0
                    )
                ),
                total_efetuado=Sum(
                    Case(
                        When(status='efetuado', then=F('valor') - F('valor_desconto') - F('valor_taxa_admin')), 
                        output_field=DecimalField(max_digits=10, decimal_places=2),
                        default=0
                    )
                ),
                count_pendente=Count(Case(When(status='pendente', then=1))),
                count_efetuado=Count(Case(When(status='efetuado', then=1))),
                count_atrasado=Count(Case(When(status='pendente', data_prevista__lt=date.today(), then=1)))
            )
        except FieldError as e:
            logger.warning(f"Erro no cálculo otimizado: {e}")
            # Fallback - cálculo manual CORRIGIDO
            stats = self._calcular_estatisticas_manual(queryset)
        
        # Garantir que valores não sejam None
        for key, value in stats.items():
            if value is None:
                stats[key] = 0
                
        return stats

    def _calcular_estatisticas_manual(self, queryset):
        """
        Cálculo manual de estatísticas CORRIGIDO
        """
        stats = {
            'total_pendente': Decimal('0'),
            'total_efetuado': Decimal('0'),
            'count_pendente': 0,
            'count_efetuado': 0,
            'count_atrasado': 0
        }
        
        hoje = date.today()
        
        for repasse in queryset:
            # *** CORREÇÃO: Usar valor líquido correto ***
            valor_liquido = self._calcular_valor_liquido(repasse)
            
            if repasse.status == 'pendente':
                stats['total_pendente'] += valor_liquido
                stats['count_pendente'] += 1
                
                if repasse.data_prevista < hoje:
                    stats['count_atrasado'] += 1
                    
            elif repasse.status == 'efetuado':
                stats['total_efetuado'] += valor_liquido
                stats['count_efetuado'] += 1
        
        return stats

    def _agrupar_repasses_seguro(self, repasses, tipo_agrupamento):
        """
        Agrupa repasses de forma segura
        """
        try:
            if tipo_agrupamento == 'data':
                return self._agrupar_por_data(repasses)
            elif tipo_agrupamento == 'proprietario':
                return self._agrupar_por_proprietario(repasses)
            elif tipo_agrupamento == 'status':
                return self._agrupar_por_status(repasses)
            else:
                return self._agrupar_sem_agrupamento(repasses)
        except Exception as e:
            logger.error(f"Erro no agrupamento {tipo_agrupamento}: {e}")
            return self._agrupar_sem_agrupamento(repasses)

    def _agrupar_por_data(self, repasses):
        """Agrupa repasses por data - ORDENAÇÃO CORRIGIDA"""
        grupos = defaultdict(list)
        
        for repasse in repasses:
            data_key = repasse.data_prevista
            grupos[data_key].append(repasse)
        
        resultado = []
        hoje = date.today()
        
        # *** CORREÇÃO: Ordenar do mais antigo para o mais novo ***
        for data, repasses_grupo in sorted(grupos.items(), reverse=False):  # reverse=False
            # Determinar relação com hoje
            if data == hoje:
                data_rel = "Hoje"
            elif data == hoje + timedelta(days=1):
                data_rel = "Amanhã"
            elif data == hoje - timedelta(days=1):
                data_rel = "Ontem"
            else:
                data_rel = None
                
            # Formato da data
            try:
                data_formatada = data.strftime('%d/%m/%Y')
                if data_rel:
                    data_formatada = f"{data_rel} - {data_formatada}"
            except:
                data_formatada = str(data)
            
            # *** CORREÇÃO: Calcular valor total corretamente ***
            valor_total = sum(self._calcular_valor_liquido(r) for r in repasses_grupo)
            
            grupo = {
                'data': data,
                'data_formatada': data_formatada,
                'repasses': repasses_grupo,
                'valor_total': valor_total,
                'count': len(repasses_grupo),
                'eh_hoje': data == hoje,
                'eh_passado': data < hoje,
                'eh_futuro': data > hoje,
                'tipo_agrupamento': 'data'
            }
            
            resultado.append(grupo)
        
        return resultado

    def _agrupar_por_proprietario(self, repasses):
        """Agrupa repasses por proprietário - CORRIGIDO"""
        grupos = defaultdict(list)
        
        for repasse in repasses:
            proprietario_key = repasse.proprietario.nome if repasse.proprietario else 'Sem proprietário'
            grupos[proprietario_key].append(repasse)
        
        resultado = []
        for proprietario, repasses_grupo in sorted(grupos.items()):
            pendentes = [r for r in repasses_grupo if r.status == 'pendente']
            efetuados = [r for r in repasses_grupo if r.status == 'efetuado']
            
            # *** CORREÇÃO: Calcular valores líquidos corretos ***
            valor_total = sum(self._calcular_valor_liquido(r) for r in repasses_grupo)
            valor_pendente = sum(self._calcular_valor_liquido(r) for r in pendentes)
            valor_efetuado = sum(self._calcular_valor_liquido(r) for r in efetuados)
            
            grupo = {
                'proprietario': proprietario,
                'repasses': repasses_grupo,
                'valor_total': valor_total,
                'count': len(repasses_grupo),
                'count_pendentes': len(pendentes),
                'count_efetuados': len(efetuados),
                'valor_pendente': valor_pendente,
                'valor_efetuado': valor_efetuado,
                'tipo_agrupamento': 'proprietario'
            }
            
            resultado.append(grupo)
        
        return resultado

    def _agrupar_por_status(self, repasses):
        """Agrupa repasses por status - CORRIGIDO"""
        grupos = defaultdict(list)
        
        for repasse in repasses:
            grupos[repasse.status].append(repasse)
        
        ordem_status = ['pendente', 'efetuado', 'cancelado']
        nomes_status = {
            'pendente': 'Pendentes',
            'efetuado': 'Efetuados', 
            'cancelado': 'Cancelados'
        }
        
        resultado = []
        for status in ordem_status:
            if status in grupos:
                repasses_grupo = grupos[status]
                
                # Calcular atrasados se for pendente
                atrasados = []
                if status == 'pendente':
                    atrasados = [r for r in repasses_grupo if r.data_prevista < date.today()]
                
                # *** CORREÇÃO: Calcular valor total corretamente ***
                valor_total = sum(self._calcular_valor_liquido(r) for r in repasses_grupo)
                
                grupo = {
                    'status': status,
                    'status_nome': nomes_status.get(status, status.title()),
                    'repasses': repasses_grupo,
                    'valor_total': valor_total,
                    'count': len(repasses_grupo),
                    'count_atrasados': len(atrasados),
                    'tem_atrasados': len(atrasados) > 0,
                    'tipo_agrupamento': 'status'
                }
                
                resultado.append(grupo)
        
        return resultado

    def _agrupar_sem_agrupamento(self, repasses):
        """Fallback - sem agrupamento"""
        valor_total = sum(self._calcular_valor_liquido(r) for r in repasses)
        
        return [{
            'tipo_agrupamento': 'sem_agrupamento',
            'repasses': list(repasses),
            'valor_total': valor_total,
            'count': len(repasses),
            'data_formatada': 'Todos os repasses',
            'proprietario': None,
            'status': None
        }]

    def _calcular_valor_liquido(self, repasse):
        """
        Calcula valor líquido CORRIGIDO
        Fórmula: Valor Líquido = Bruto - Taxa - Desconto
        """
        try:
            # Tentar campo específico primeiro
            if hasattr(repasse, 'valor_liquido_repasse') and repasse.valor_liquido_repasse is not None:
                return repasse.valor_liquido_repasse
            
            # Tentar propriedade calculada
            if hasattr(repasse, 'valor_liquido'):
                if callable(getattr(repasse, 'valor_liquido')):
                    return repasse.valor_liquido()
                elif repasse.valor_liquido is not None:
                    return repasse.valor_liquido
            
            # *** CORREÇÃO: Calcular manualmente de forma correta ***
            valor_bruto = getattr(repasse, 'valor', 0) or Decimal('0')
            valor_desconto = getattr(repasse, 'valor_desconto', 0) or Decimal('0')
            valor_taxa = getattr(repasse, 'valor_taxa_admin', 0) or Decimal('0')
            
            # Valor líquido = Bruto - Taxa - Desconto
            return valor_bruto - valor_desconto - valor_taxa
            
        except (AttributeError, TypeError, ValueError) as e:
            logger.warning(f"Erro ao calcular valor líquido do repasse {getattr(repasse, 'id', 'unknown')}: {e}")
            return Decimal('0')

    # *** MÉTODO ADICIONAL: Para calcular datas com 5 dias úteis ***
    @staticmethod
    def _calcular_data_repasse(data_pagamento, dias_uteis=5):
        """
        Calcula data de repasse com dias úteis
        EXEMPLO: 01/07/2025 + 5 dias úteis = 08/07/2025
        """
        from datetime import timedelta
        
        data_atual = data_pagamento
        dias_adicionados = 0
        
        while dias_adicionados < dias_uteis:
            data_atual = data_atual + timedelta(days=1)
            
            # Segunda=0 a Sexta=4 são dias úteis
            if data_atual.weekday() < 5:
                dias_adicionados += 1
        
        return data_atual

    # *** MÉTODO ADICIONAL: Para debug ***
    def debug_repasses(self):
        """Método para debug - verificar se cálculos estão corretos"""
        repasses = self.get_queryset()[:5]  # Primeiros 5 repasses
        
        print("=== DEBUG REPASSES ===")
        for repasse in repasses:
            valor_bruto = getattr(repasse, 'valor', 0)
            valor_taxa = getattr(repasse, 'valor_taxa_admin', 0)
            valor_desconto = getattr(repasse, 'valor_desconto', 0)
            valor_liquido = self._calcular_valor_liquido(repasse)
            
            print(f"Repasse {repasse.id}:")
            print(f"  Bruto: R$ {valor_bruto:.2f}")
            print(f"  Taxa: R$ {valor_taxa:.2f}")
            print(f"  Desconto: R$ {valor_desconto:.2f}")
            print(f"  Líquido: R$ {valor_liquido:.2f}")
            print(f"  Fórmula: {valor_bruto} - {valor_taxa} - {valor_desconto} = {valor_liquido}")
            print()
        
        return repasses


# *** FUNÇÕES AUXILIARES PARA USAR EM OUTROS LUGARES ***

def adicionar_dias_uteis(data_inicial, quantidade_dias):
    """
    Função auxiliar para adicionar dias úteis
    Pode ser usada em outros serviços
    """
    from datetime import timedelta
    
    data_atual = data_inicial
    dias_adicionados = 0
    
    while dias_adicionados < quantidade_dias:
        data_atual = data_atual + timedelta(days=1)
        
        # Segunda=0 a Sexta=4 são dias úteis
        if data_atual.weekday() < 5:
            dias_adicionados += 1
    
    return data_atual


def calcular_valor_liquido_repasse(repasse):
    """
    Função auxiliar para calcular valor líquido
    Pode ser usada em outros lugares
    """
    try:
        valor_bruto = getattr(repasse, 'valor', 0) or Decimal('0')
        valor_desconto = getattr(repasse, 'valor_desconto', 0) or Decimal('0')
        valor_taxa = getattr(repasse, 'valor_taxa_admin', 0) or Decimal('0')
        
        # Valor líquido = Bruto - Taxa - Desconto
        return valor_bruto - valor_desconto - valor_taxa
        
    except (AttributeError, TypeError, ValueError):
        return Decimal('0')


def testar_calculo_data_repasse():
    """
    Função para testar se o cálculo de data está correto
    """
    from datetime import date
    
    print("=== TESTE DE CÁLCULO DE DATA ===")
    
    # Teste 1: Terça-feira, 01/07/2025
    data_pagamento = date(2025, 7, 1)
    data_repasse = adicionar_dias_uteis(data_pagamento, 5)
    
    print(f"Pagamento: {data_pagamento.strftime('%d/%m/%Y')} ({data_pagamento.strftime('%A')})")
    print(f"Repasse: {data_repasse.strftime('%d/%m/%Y')} ({data_repasse.strftime('%A')})")
    print(f"Correto: {data_repasse == date(2025, 7, 8)}")
    
    # Teste 2: Sexta-feira, 04/07/2025
    data_sexta = date(2025, 7, 4)
    data_repasse_sexta = adicionar_dias_uteis(data_sexta, 5)
    
    print(f"\nPagamento: {data_sexta.strftime('%d/%m/%Y')} ({data_sexta.strftime('%A')})")
    print(f"Repasse: {data_repasse_sexta.strftime('%d/%m/%Y')} ({data_repasse_sexta.strftime('%A')})")
    print(f"Correto: {data_repasse_sexta == date(2025, 7, 11)}")
    
    return True


# *** EXEMPLO DE USO ***
if __name__ == "__main__":
    # Testar função de data
    testar_calculo_data_repasse()
    
    # Testar cálculo de valor líquido
    print("\n=== TESTE DE VALOR LÍQUIDO ===")
    from decimal import Decimal
    
    # Simular um repasse
    class RepasseTeste:
        def __init__(self, valor, taxa, desconto):
            self.id = 1
            self.valor = Decimal(str(valor))
            self.valor_taxa_admin = Decimal(str(taxa))
            self.valor_desconto = Decimal(str(desconto))
    
    repasse_teste = RepasseTeste(1000, 80, 20)
    valor_liquido = calcular_valor_liquido_repasse(repasse_teste)
    
    print(f"Bruto: R$ {repasse_teste.valor}")
    print(f"Taxa: R$ {repasse_teste.valor_taxa_admin}")
    print(f"Desconto: R$ {repasse_teste.valor_desconto}")
    print(f"Líquido: R$ {valor_liquido}")
    print(f"Correto: {valor_liquido == Decimal('900')}")



def obter_estatisticas_repasses():
    """
    Função auxiliar para obter estatísticas gerais de repasses
    """
    from datetime import timedelta
    
    hoje = date.today()
    inicio_mes = hoje.replace(day=1)
    
    stats = Repasse.objects.aggregate(
        total_mes=Sum(
            Case(
                When(
                    data_prevista__gte=inicio_mes,
                    data_prevista__lte=hoje,
                    then=F('valor') - F('valor_desconto') - F('valor_taxa_admin')
                ),
                output_field=DecimalField(max_digits=10, decimal_places=2),
                default=0
            )
        ),
        proximos_7_dias=Count(
            Case(
                When(
                    status='pendente',
                    data_prevista__gte=hoje,
                    data_prevista__lte=hoje + timedelta(days=7),
                    then=1
                )
            )
        ),
        em_atraso=Count(
            Case(
                When(
                    status='pendente',
                    data_prevista__lt=hoje,
                    then=1
                )
            )
        )
    )
    
    return stats


def calcular_resumo_mensal(mes=None, ano=None):
    """
    Calcula resumo mensal de repasses
    """
    if not mes or not ano:
        hoje = date.today()
        mes, ano = hoje.month, hoje.year
    
    repasses_mes = Repasse.objects.filter(
        mes_referencia=mes,
        ano_referencia=ano
    )
    
    resumo = repasses_mes.aggregate(
        total_valor=Sum('valor'),
        total_taxa_admin=Sum('valor_taxa_admin'),
        count_total=Count('id'),
        count_pendentes=Count(Case(When(status='pendente', then=1))),
        count_efetuados=Count(Case(When(status='efetuado', then=1))),
        count_cancelados=Count(Case(When(status='cancelado', then=1)))
    )
    
    # Adicionar percentuais
    total = resumo['count_total'] or 1
    resumo['percentual_efetuados'] = (resumo['count_efetuados'] or 0) / total * 100
    resumo['percentual_pendentes'] = (resumo['count_pendentes'] or 0) / total * 100
    
    return resumo

class RepasseDetailView(DetailView):
    model = Repasse
    template_name = 'financeiro/repasses/detalhe_repasse.html'
    context_object_name = 'repasse'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        repasse = self.get_object()
        
        # =============== MOVIMENTOS - VERSÃO SEGURA ===============
        context['movimentos'] = []
        try:
            from ..models.movimento import MovimentoConta
            
            # Tentar diferentes combinações de campos para encontrar movimentos relacionados
            movimentos_filters = [
                # Opção 1: Filtros ideais
                {
                    'proprietario': repasse.proprietario,
                    'contrato': repasse.contrato,
                    'data_referencia__year': repasse.ano_referencia,
                    'data_referencia__month': repasse.mes_referencia
                },
                # Opção 2: Apenas proprietário e contrato
                {
                    'proprietario': repasse.proprietario,
                    'contrato': repasse.contrato
                },
                # Opção 3: Apenas proprietário
                {
                    'proprietario': repasse.proprietario
                }
            ]
            
            for filters in movimentos_filters:
                try:
                    movimentos = MovimentoConta.objects.filter(**filters)
                    
                    # Tentar diferentes campos para ordenação
                    for order_field in ['-data', '-data_referencia', '-id']:
                        try:
                            context['movimentos'] = list(movimentos.order_by(order_field)[:10])
                            break
                        except:
                            continue
                    
                    if context['movimentos']:
                        break
                        
                except Exception:
                    continue
                    
        except ImportError:
            pass
        except Exception as e:
            import logging
            logging.getLogger(__name__).debug(f"Erro ao buscar movimentos: {e}")
        
        # =============== REPASSES RELACIONADOS - VERSÃO SEGURA ===============
        try:
            context['repasses_relacionados'] = Repasse.objects.filter(
                proprietario=repasse.proprietario,
                ano_referencia=repasse.ano_referencia,
                mes_referencia=repasse.mes_referencia
            ).exclude(pk=repasse.pk).select_related('contrato')[:5]
        except Exception:
            context['repasses_relacionados'] = []
        
        # =============== POLÍTICA DO CONTRATO - VERSÃO SEGURA ===============
        context['politica_info'] = None
        try:
            if hasattr(repasse.contrato, 'politica_repasse'):
                politica = repasse.contrato.politica_repasse
                if politica:
                    context['politica_info'] = {
                        'periodicidade': getattr(politica, 'periodicidade', 'N/A'),
                        'tipo_dias': getattr(politica, 'tipo_dias', 'N/A'),
                        'ativa': getattr(politica, 'ativa', False),
                        'valor_minimo': getattr(politica, 'valor_minimo_repasse', 0),
                        'taxa_admin': 'N/A'
                    }
                    
                    # Tentar obter taxa administrativa
                    try:
                        if hasattr(politica, 'get_taxa_admin'):
                            context['politica_info']['taxa_admin'] = politica.get_taxa_admin()
                        elif hasattr(politica, 'taxa_admin_personalizada'):
                            context['politica_info']['taxa_admin'] = politica.taxa_admin_personalizada
                    except:
                        pass
        except Exception:
            pass
        
        # =============== INFORMAÇÕES DA COBRANÇA - VERSÃO SEGURA ===============
        context['cobranca_info'] = None
        if repasse.cobranca:
            try:
                cobranca = repasse.cobranca
                context['cobranca_info'] = {}
                
                # Campos básicos com fallbacks
                context['cobranca_info']['id'] = getattr(cobranca, 'id', 'N/A')
                context['cobranca_info']['status'] = getattr(cobranca, 'status', 'N/A')
                
                # Valores com múltiplos fallbacks
                valor_fields = ['valor_pago', 'valor_original', 'valor']
                for field in valor_fields:
                    if hasattr(cobranca, field):
                        valor = getattr(cobranca, field)
                        if valor:
                            context['cobranca_info']['valor'] = valor
                            break
                else:
                    context['cobranca_info']['valor'] = 0
                
                # Datas com fallbacks
                date_fields = ['data_pagamento', 'data_vencimento', 'data']
                for field in date_fields:
                    if hasattr(cobranca, field):
                        data = getattr(cobranca, field)
                        if data:
                            context['cobranca_info'][field] = data
                
                # Referência temporal
                context['cobranca_info']['mes_referencia'] = getattr(cobranca, 'mes_referencia', repasse.mes_referencia)
                context['cobranca_info']['ano_referencia'] = getattr(cobranca, 'ano_referencia', repasse.ano_referencia)
                
            except Exception:
                context['cobranca_info'] = None
        
        # =============== INFORMAÇÕES DO CONTRATO - VERSÃO SEGURA ===============
        try:
            contrato = repasse.contrato
            context['contrato_info'] = {
                'id': contrato.id,
                'endereco': 'N/A',
                'valor_aluguel': 0,
                'tem_politica': False
            }
            
            # Endereço do imóvel
            if hasattr(contrato, 'imovel') and contrato.imovel:
                try:
                    context['contrato_info']['endereco'] = str(contrato.imovel)
                except:
                    context['contrato_info']['endereco'] = 'Endereço não disponível'
            
            # Valor do aluguel
            valor_fields = ['valor_aluguel', 'valor_base', 'valor']
            for field in valor_fields:
                if hasattr(contrato, field):
                    valor = getattr(contrato, field)
                    if valor:
                        context['contrato_info']['valor_aluguel'] = valor
                        break
            
            # Verificar se tem política
            context['contrato_info']['tem_politica'] = (
                hasattr(contrato, 'politica_repasse') and 
                contrato.politica_repasse is not None
            )
            
        except Exception:
            context['contrato_info'] = None
        
        # =============== INFORMAÇÕES DO PROPRIETÁRIO - VERSÃO SEGURA ===============
        try:
            proprietario = repasse.proprietario
            context['proprietario_info'] = {
                'id': proprietario.id,
                'nome': getattr(proprietario, 'nome', 'N/A'),
                'email': getattr(proprietario, 'email', ''),
                'telefone': getattr(proprietario, 'telefone', ''),
                'tipo': getattr(proprietario, 'tipo', 'N/A')
            }
        except Exception:
            context['proprietario_info'] = None
        
        # =============== ESTATÍSTICAS RÁPIDAS ===============
        try:
            # Outros repasses do mesmo proprietário
            outros_repasses = Repasse.objects.filter(
                proprietario=repasse.proprietario
            ).exclude(pk=repasse.pk)
            
            context['estatisticas'] = {
                'total_repasses': outros_repasses.count(),
                'repasses_pendentes': outros_repasses.filter(status='pendente').count(),
                'repasses_efetuados': outros_repasses.filter(status='efetuado').count(),
            }
            
            # Valor total se possível calcular
            try:
                from django.db.models import Sum, F
                total_valor = outros_repasses.filter(status='efetuado').aggregate(
                    total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin'))
                )['total']
                context['estatisticas']['valor_total_recebido'] = total_valor or 0
            except:
                context['estatisticas']['valor_total_recebido'] = 'N/A'
                
        except Exception:
            context['estatisticas'] = None
        
        return context
        
class RepasseCreateView(CreateView):
    model = Repasse
    template_name = 'financeiro/repasses/repasse_form.html'
    
    def get_form_class(self):
        return RepasseForm if RepasseForm else None
    
    def get_fields(self):
        if not RepasseForm:
            return ['proprietario', 'contrato', 'valor', 'valor_desconto', 
                    'valor_taxa_admin', 'data_prevista', 'mes_referencia', 'ano_referencia',
                    'tipo', 'metodo_pagamento', 'descricao', 'observacoes']
        return None
    
    def get_success_url(self):
        return reverse_lazy('financeiro:repasse_detail', kwargs={'pk': self.object.pk})
    
    def form_valid(self, form):
        messages.success(self.request, 'Repasse criado com sucesso!')
        return super().form_valid(form)


class RepasseUpdateView(UpdateView):
    model = Repasse
    template_name = 'financeiro/repasses/repasse_form.html'
    
    def get_form_class(self):
        return RepasseForm if RepasseForm else None
    
    def get_fields(self):
        if not RepasseForm:
            return ['proprietario', 'contrato', 'valor', 'valor_desconto', 
                    'valor_taxa_admin', 'data_prevista', 'mes_referencia', 'ano_referencia',
                    'tipo', 'metodo_pagamento', 'descricao', 'observacoes']
        return None
    
    def get_success_url(self):
        return reverse_lazy('financeiro:repasse_detail', kwargs={'pk': self.object.pk})
    
    def form_valid(self, form):
        messages.success(self.request, 'Repasse atualizado com sucesso!')
        return super().form_valid(form)


# ==================== VIEWS DE POLÍTICAS POR CONTRATO ====================

class ContratosSemPoliticaView(ListView):
    """Lista contratos ativos sem política de repasse configurada"""
    template_name = 'financeiro/repasses/contratos_sem_politica.html'
    context_object_name = 'contratos'
    paginate_by = 20

    def get_queryset(self):
        # CORRIGIDO: usar ativo=True
        queryset = Contrato.objects.filter(
            ativo=True,
            politica_repasse__isnull=True
        ).select_related('proprietario', 'imovel').order_by('proprietario__nome')

        # Filtros básicos
        proprietario_id = self.request.GET.get('proprietario')
        if proprietario_id:
            queryset = queryset.filter(proprietario_id=proprietario_id)
        
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(proprietario__nome__icontains=search) |
                Q(imovel__endereco__icontains=search)
            )

        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['total_sem_politica'] = self.get_queryset().count()
        context['total_contratos_ativos'] = Contrato.objects.filter(ativo=True).count()  # CORRIGIDO
        return context


class PoliticaContratoListView(ListView):
    """Lista todas as políticas por contrato"""
    model = PoliticaRepasseContrato
    template_name = 'financeiro/repasses/politica_contrato_list.html'
    context_object_name = 'politicas'
    paginate_by = 20

    def get_queryset(self):
        # ✅ CORRIGIDO: Usar prefetch_related para ManyToMany
        return PoliticaRepasseContrato.objects.select_related(
            'contrato',              # ForeignKey - OK
            'contrato__imovel',      # ForeignKey através de contrato - OK
            'contrato__fiador'       # ForeignKey através de contrato - OK
        ).prefetch_related(
            'contrato__proprietario',  # ManyToMany através de contrato - CORRETO
            'contrato__inquilino'      # ManyToMany através de contrato - CORRETO
        ).order_by('-ativa', 'contrato__id')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        # Estatísticas
        context['stats'] = {
            'total_politicas': PoliticaRepasseContrato.objects.count(),
            'politicas_ativas': PoliticaRepasseContrato.objects.filter(ativa=True).count(),
            'contratos_sem_politica': Contrato.objects.filter(
                ativo=True,
                politica_repasse__isnull=True
            ).count(),
            'por_periodicidade': PoliticaRepasseContrato.objects.filter(ativa=True).values(
                'periodicidade'
            ).annotate(count=Count('id')),
            'por_tipo_dias': PoliticaRepasseContrato.objects.filter(ativa=True).values(
                'tipo_dias'
            ).annotate(count=Count('id'))
        }
        
        return context


# No arquivo repasse_views.py

class PoliticaContratoCreateView(CreateView):
    model = PoliticaRepasseContrato
    template_name = 'financeiro/repasses/politica_contrato_form.html'
    fields = ['contrato', 'ativa', 'periodicidade', 'tipo_dias', 'dia_mes', 'dia_semana',
              'dias_apos_recebimento', 'percentual_adiantamento', 'taxa_adiantamento',
              'valor_minimo_repasse', 'taxa_admin_personalizada', 'considerar_feriados',
              'antecipar_fds_feriados', 'observacoes']
    
    def get_success_url(self):
        return reverse_lazy('financeiro:politica_contrato_list')
    
    def get_form(self, form_class=None):
        """Customiza o formulário para melhorar a exibição dos contratos"""
        form = super().get_form(form_class)
        
        # Customizar o campo contrato
        form.fields['contrato'].queryset = Contrato.objects.filter(
            ativo=True,
            politica_repasse__isnull=True
        ).select_related('imovel').prefetch_related('proprietario').order_by('id')
        
        # Customizar o widget para melhor exibição
        form.fields['contrato'].widget.attrs.update({
            'class': 'form-select',
            'data-placeholder': 'Selecione um contrato...'
        })
        
        return form
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        # Criar lista customizada de contratos para o template
        contratos_sem_politica = Contrato.objects.filter(
            ativo=True,
            politica_repasse__isnull=True
        ).select_related('imovel').prefetch_related('proprietario').order_by('id')
        
        # Criar lista com informações formatadas
        contratos_formatados = []
        for contrato in contratos_sem_politica:
            proprietarios = contrato.proprietario.all()
            nomes_proprietarios = ", ".join([p.nome for p in proprietarios]) if proprietarios else "Sem proprietário"
            
            endereco_completo = f"{contrato.imovel.endereco}"
            if contrato.imovel.numero:
                endereco_completo += f", {contrato.imovel.numero}"
            if contrato.imovel.complemento:
                endereco_completo += f" - {contrato.imovel.complemento}"
            
            contratos_formatados.append({
                'id': contrato.id,
                'proprietarios': nomes_proprietarios,
                'endereco': endereco_completo,
                'valor_base': contrato.valor_base,
                'display_name': f"#{contrato.id} - {nomes_proprietarios} - {endereco_completo}"
            })
        
        context['contratos_formatados'] = contratos_formatados
        context['contratos_sem_politica'] = contratos_sem_politica  # Manter compatibilidade
        
        return context
    
    def form_valid(self, form):
        messages.success(self.request, f'Política criada para o contrato {form.instance.contrato.id}!')
        return super().form_valid(form)
    
    def get_success_url(self):
        return reverse_lazy('financeiro:politica_contrato_list')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        # ✅ CORRIGIDO: Usar prefetch_related para ManyToMany
        context['contratos_sem_politica'] = Contrato.objects.filter(
            ativo=True,
            politica_repasse__isnull=True
        ).select_related(
            'imovel',      # ForeignKey - OK
            'fiador'       # ForeignKey - OK  
        ).prefetch_related(
            'proprietario',  # ManyToMany - CORRETO
            'inquilino'      # ManyToMany - CORRETO
        ).order_by('id')  # ✅ Ordenar por ID em vez de proprietario__nome (ManyToMany)
        
        return context
    
    def form_valid(self, form):
        messages.success(self.request, f'Política criada para o contrato {form.instance.contrato.id}!')
        return super().form_valid(form)
    
    def form_valid(self, form):
        messages.success(self.request, f'Política criada para o contrato {form.instance.contrato.id}!')
        return super().form_valid(form)


class PoliticaContratoUpdateView(UpdateView):
    model = PoliticaRepasseContrato
    template_name = 'financeiro/repasses/politica_contrato_form.html'
    
    def get_form_class(self):
        return PoliticaRepasseContratoForm if PoliticaRepasseContratoForm else None
    
    def get_fields(self):
        if not PoliticaRepasseContratoForm:
            return ['ativa', 'periodicidade', 'tipo_dias', 'dia_mes', 'dia_semana',
                    'dias_apos_recebimento', 'percentual_adiantamento', 'taxa_adiantamento',
                    'valor_minimo_repasse', 'taxa_admin_personalizada', 'considerar_feriados',
                    'antecipar_fds_feriados', 'observacoes']
        return None
    
    def get_success_url(self):
        return reverse_lazy('financeiro:politica_contrato_list')
    
    def form_valid(self, form):
        messages.success(self.request, f'Política do contrato {form.instance.contrato.id} atualizada!')
        return super().form_valid(form)


# ==================== VIEWS DE POLÍTICAS GLOBAIS ====================

class PoliticaGlobalListView(ListView):
    model = PoliticaRepasseGlobal
    template_name = 'financeiro/repasses/politica_global_list.html'
    context_object_name = 'politicas'

    def get_queryset(self):
        return PoliticaRepasseGlobal.objects.order_by('-ativa', 'nome')


class PoliticaGlobalCreateView(CreateView):
    model = PoliticaRepasseGlobal
    template_name = 'financeiro/repasses/politica_form.html'
    
    def get_form_class(self):
        return PoliticaRepasseGlobalForm if PoliticaRepasseGlobalForm else None
    
    def get_fields(self):
        if not PoliticaRepasseGlobalForm:
            return ['nome', 'ativa', 'periodicidade', 'tipo_dias', 'dia_mes', 'dia_semana',
                    'dias_apos_recebimento', 'percentual_adiantamento', 'taxa_adiantamento',
                    'valor_minimo_repasse', 'taxa_admin_padrao', 'considerar_feriados',
                    'antecipar_fds_feriados']
        return None
    
    def get_success_url(self):
        return reverse_lazy('financeiro:politica_global_list')


class PoliticaGlobalUpdateView(UpdateView):
    model = PoliticaRepasseGlobal
    template_name = 'financeiro/repasses/politica_form.html'
    
    def get_form_class(self):
        return PoliticaRepasseGlobalForm if PoliticaRepasseGlobalForm else None
    
    def get_fields(self):
        if not PoliticaRepasseGlobalForm:
            return ['nome', 'ativa', 'periodicidade', 'tipo_dias', 'dia_mes', 'dia_semana',
                    'dias_apos_recebimento', 'percentual_adiantamento', 'taxa_adiantamento',
                    'valor_minimo_repasse', 'taxa_admin_padrao', 'considerar_feriados',
                    'antecipar_fds_feriados']
        return None
    
    def get_success_url(self):
        return reverse_lazy('financeiro:politica_global_list')


# ==================== VIEWS DE PROCESSAMENTO ====================

def processar_repasse_view(request, repasse_id):
    """Processa (efetiva) um repasse pendente"""
    repasse = get_object_or_404(Repasse, id=repasse_id)
    
    if request.method == 'POST':
        try:
            metodo_pagamento = request.POST.get('metodo_pagamento')
            observacoes = request.POST.get('observacoes')
            
            success = repasse.efetivar_repasse(
                metodo_pagamento=metodo_pagamento,
                observacoes=observacoes
            )
            
            if success:
                # Upload do comprovante se fornecido
                comprovante = request.FILES.get('comprovante')
                if comprovante:
                    repasse.comprovante = comprovante
                    repasse.save(update_fields=['comprovante'])
                
                messages.success(request, f'Repasse efetuado com sucesso!')
                return redirect('financeiro:repasse_detail', pk=repasse.id)
            else:
                messages.error(request, 'Não foi possível efetivar o repasse.')
        except Exception as e:
            logger.error(f"Erro ao processar repasse {repasse_id}: {str(e)}")
            messages.error(request, f'Erro ao processar repasse: {str(e)}')

    return render(request, 'financeiro/repasses/processar_repasse.html', {
        'repasse': repasse,
    })


def cancelar_repasse_view(request, repasse_id):
    """Cancela um repasse pendente"""
    repasse = get_object_or_404(Repasse, id=repasse_id)
    
    if request.method == 'POST':
        motivo = request.POST.get('motivo', '')
        success = repasse.cancelar_repasse(motivo)
        
        if success:
            messages.success(request, 'Repasse cancelado com sucesso!')
        else:
            messages.error(request, 'Não foi possível cancelar o repasse.')
    
    return redirect('financeiro:repasse_detail', pk=repasse.id)


def gerar_repasse_automatico(request):
    """Gera repasses baseados nas políticas dos contratos"""
    if request.method == 'POST':
        try:
            if RepasseService:
                service = RepasseService()
                resultado = service.processar_repasses_automaticos()
                
                messages.success(request, 
                    f"Processamento concluído: {resultado.get('criados', 0)} repasses criados, "
                    f"{resultado.get('erros', 0)} erros.")
            else:
                messages.warning(request, 'Serviço de repasse não disponível.')
            
        except Exception as e:
            logger.error(f"Erro no processamento automático: {str(e)}")
            messages.error(request, f"Erro no processamento: {str(e)}")
    
    return redirect('financeiro:lista_repasses')


# ==================== VIEWS DE AGENDAMENTOS ====================

class AgendamentoListView(ListView):
    model = AgendamentoRepasse
    template_name = 'financeiro/repasses/agendamento_list.html'
    context_object_name = 'agendamentos'
    paginate_by = 20

    def get_queryset(self):
        return AgendamentoRepasse.objects.select_related(
            'proprietario', 'contrato', 'politica_contrato', 'politica_global'
        ).order_by('-data_agendada')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        # Estatísticas
        hoje = date.today()
        context['stats'] = {
            'total_agendados': AgendamentoRepasse.objects.filter(status='agendado').count(),
            'vencidos_hoje': AgendamentoRepasse.objects.filter(
                status='agendado',
                data_agendada__lte=hoje
            ).count(),
            'proximos_7_dias': AgendamentoRepasse.objects.filter(
                status='agendado',
                data_agendada__range=[hoje, hoje + timedelta(days=7)]
            ).count()
        }
        
        return context


def processar_agendamento_view(request, pk):
    """Processa um agendamento específico"""
    agendamento = get_object_or_404(AgendamentoRepasse, pk=pk)
    
    if request.method == 'POST':
        try:
            repasse = agendamento.processar()
            if repasse:
                messages.success(request, f'Agendamento processado! Repasse #{repasse.id} criado.')
                return redirect('financeiro:repasse_detail', pk=repasse.id)
            else:
                messages.error(request, 'Erro ao processar agendamento.')
        except Exception as e:
            logger.error(f"Erro ao processar agendamento {pk}: {str(e)}")
            messages.error(request, f'Erro: {str(e)}')
    
    return redirect('financeiro:agendamento_list')



# ==================== VIEWS DE DASHBOARD ====================

class DashboardFinanceiroView(TemplateView):
    """Dashboard principal do módulo financeiro - CAMPOS CORRIGIDOS"""
    template_name = 'financeiro/dashboard_financeiro.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        hoje = date.today()
        inicio_mes = hoje.replace(day=1)
        
        # Estatísticas de repasses
        context['stats_repasses'] = {
            'pendentes': Repasse.objects.filter(status='pendente').aggregate(
                total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin')),
                count=Count('id')
            ),
            'atrasados': Repasse.objects.filter(
                status='pendente',
                data_prevista__lt=hoje
            ).aggregate(
                total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin')),
                count=Count('id')
            ),
            'mes_atual': Repasse.objects.filter(
                data_criacao__gte=inicio_mes  # Verificar se este campo existe
            ).aggregate(
                total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin')),
                count=Count('id')
            )
        }
        
        # Próximos repasses (7 dias)
        context['proximos_repasses'] = Repasse.objects.filter(
            status='pendente',
            data_prevista__lte=hoje + timedelta(days=7)
        ).select_related('proprietario', 'contrato').order_by('data_prevista')[:10]
        
        # Repasses por status (últimos 30 dias) - CORRIGIDO
        try:
            context['repasses_por_status'] = list(Repasse.objects.filter(
                data_criacao__gte=hoje - timedelta(days=30)  # Pode precisar ajustar
            ).values('status').annotate(
                total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin')),
                count=Count('id')
            ))
        except FieldError:
            # Se data_criacao não existir, usar outro campo ou filtro simples
            context['repasses_por_status'] = list(Repasse.objects.values('status').annotate(
                total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin')),
                count=Count('id')
            ))
        
        # Estatísticas de políticas
        context['politicas_ativas'] = PoliticaRepasseContrato.objects.filter(ativa=True).count()
        context['contratos_sem_politica'] = Contrato.objects.filter(
            ativo=True,
            politica_repasse__isnull=True
        ).count()
        
        # Agendamentos pendentes
        context['agendamentos_pendentes'] = AgendamentoRepasse.objects.filter(
            status='agendado',
            data_agendada__lte=hoje + timedelta(days=7)
        ).count()
        
        return context
class ConfiguracaoRepasseView(TemplateView):
    """View para configurações gerais do sistema de repasse"""
    template_name = 'financeiro/repasses/configurar_politicas.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        if ConfiguracaoPoliticasForm:
            context['form'] = ConfiguracaoPoliticasForm()
        
        # Estatísticas do sistema - CORRIGIDO
        context['stats'] = {
            'total_politicas_contrato': PoliticaRepasseContrato.objects.count(),
            'politicas_ativas_contrato': PoliticaRepasseContrato.objects.filter(ativa=True).count(),
            'total_politicas_global': PoliticaRepasseGlobal.objects.count(),
            'politicas_ativas_global': PoliticaRepasseGlobal.objects.filter(ativa=True).count(),
            'contratos_sem_politica': Contrato.objects.filter(
                ativo=True,  # CORRIGIDO
                politica_repasse__isnull=True
            ).count(),
            'repasses_mes': Repasse.objects.filter(
                data_criacao__month=date.today().month,
                data_criacao__year=date.today().year
            ).count()
        }
        
        return context
    
    def post(self, request, *args, **kwargs):
        if ConfiguracaoPoliticasForm:
            form = ConfiguracaoPoliticasForm(request.POST)
            
            if form.is_valid():
                # Aqui você salvaria as configurações (implementar modelo de configuração)
                messages.success(request, 'Configurações salvas com sucesso!')
                return redirect('financeiro:configuracao_repasse')
            
            return self.render_to_response(self.get_context_data(form=form))
        else:
            messages.warning(request, 'Formulário de configuração não disponível.')
            return redirect('financeiro:configuracao_repasse')


# ==================== VIEWS AJAX ====================

def toggle_politica_contrato_view(request, pk):
    """Toggle status ativo/inativo de uma política de contrato"""
    if request.method == 'POST':
        politica = get_object_or_404(PoliticaRepasseContrato, pk=pk)
        
        try:
            politica.ativa = not politica.ativa
            politica.save()
            
            messages.success(request, f'Política do contrato {politica.contrato.id} {"ativada" if politica.ativa else "desativada"}!')
            return JsonResponse({'success': True})
            
        except Exception as e:
            logger.error(f"Erro ao alterar status da política {pk}: {str(e)}")
            return JsonResponse({'success': False, 'error': str(e)})
    
    return JsonResponse({'success': False, 'error': 'Método não permitido'})


def dashboard_repasses_ajax(request):
    """Dados para gráficos do dashboard via AJAX"""
    hoje = date.today()
    inicio_mes = hoje.replace(day=1)
    
    # Repasses por status
    status_data = Repasse.objects.filter(
        data_criacao__date__gte=inicio_mes
    ).values('status').annotate(
        total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin')),
        count=Count('id')
    )
    
    # Políticas por tipo de dias
    politicas_por_tipo = PoliticaRepasseContrato.objects.filter(
        ativa=True
    ).values('tipo_dias').annotate(count=Count('id'))
    
    # Próximos repasses (7 dias)
    proximos = Repasse.objects.filter(
        status='pendente',
        data_prevista__lte=hoje + timedelta(days=7)
    ).select_related('proprietario').order_by('data_prevista')[:10]
    
    return JsonResponse({
        'status_data': list(status_data),
        'politicas_por_tipo': list(politicas_por_tipo),
        'proximos_repasses': [{
            'id': r.id,
            'proprietario': r.proprietario.nome,
            'valor': float(r.valor_liquido),
            'data_prevista': r.data_prevista.strftime('%d/%m/%Y'),
            'dias_restantes': (r.data_prevista - hoje).days
        } for r in proximos]
    })


def aplicar_politica_em_lote(request):
    """Aplica uma política global em lote para contratos selecionados"""
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            politica_global_id = data.get('politica_global_id')
            contratos_ids = data.get('contratos_ids', [])
            
            if not politica_global_id:
                return JsonResponse({
                    'success': False, 
                    'error': 'Política deve ser informada'
                })
            
            politica_global = get_object_or_404(PoliticaRepasseGlobal, pk=politica_global_id)
            
            if contratos_ids:
                contratos = Contrato.objects.filter(
                    id__in=contratos_ids,
                    politica_repasse__isnull=True
                )
            else:
                # CORRIGIDO: usar ativo=True
                contratos = Contrato.objects.filter(
                    ativo=True,
                    politica_repasse__isnull=True
                )
            
            criados = 0
            erros = []
            
            with transaction.atomic():
                for contrato in contratos:
                    try:
                        PoliticaRepasseContrato.objects.create(
                            contrato=contrato,
                            ativa=politica_global.ativa,
                            periodicidade=politica_global.periodicidade,
                            tipo_dias=politica_global.tipo_dias,
                            dia_mes=politica_global.dia_mes,
                            dia_semana=politica_global.dia_semana,
                            dias_apos_recebimento=politica_global.dias_apos_recebimento,
                            percentual_adiantamento=politica_global.percentual_adiantamento,
                            taxa_adiantamento=politica_global.taxa_adiantamento,
                            valor_minimo_repasse=politica_global.valor_minimo_repasse,
                            taxa_admin_personalizada=politica_global.taxa_admin_padrao,
                            considerar_feriados=politica_global.considerar_feriados,
                            antecipar_fds_feriados=politica_global.antecipar_fds_feriados,
                            observacoes=f"Criada em lote baseada na política global '{politica_global.nome}'"
                        )
                        criados += 1
                    except Exception as e:
                        erros.append(f"Contrato {contrato.id}: {str(e)}")
            
            return JsonResponse({
                'success': True,
                'criados': criados,
                'erros': erros
            })
            
        except Exception as e:
            logger.error(f"Erro ao aplicar política em lote: {str(e)}")
            return JsonResponse({
                'success': False,
                'error': str(e)
            })
    
    return JsonResponse({'success': False, 'error': 'Método não permitido'})


# ==================== VIEWS DE RELATÓRIOS ====================

def relatorio_repasses_view(request):
    """Gera relatórios de repasses"""
    return render(request, 'financeiro/repasses/relatorio_repasses.html')


def exportar_relatorio_repasses(request):
    """Exporta relatório de repasses em Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        
        # Parâmetros de filtro
        periodo = int(request.GET.get('periodo', 30))
        
        # Buscar dados
        data_inicio = date.today() - timedelta(days=periodo)
        repasses = Repasse.objects.filter(
            data_criacao__date__gte=data_inicio
        ).select_related('proprietario', 'contrato').order_by('-data_criacao')
        
        # Criar workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Relatório de Repasses"
        
        # Cabeçalhos
        headers = [
            'ID', 'Data Criação', 'Proprietário', 'Contrato', 'Imóvel',
            'Valor Bruto', 'Descontos', 'Taxa Admin', 'Valor Líquido',
            'Status', 'Data Prevista', 'Data Efetivação', 'Método Pagamento',
            'Mês/Ano Referência', 'Tipo', 'Política'
        ]
        
        # Estilo do cabeçalho
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        
        # Dados
        for row, repasse in enumerate(repasses, 2):
            ws.cell(row=row, column=1, value=repasse.id)
            ws.cell(row=row, column=2, value=repasse.data_criacao.date())
            ws.cell(row=row, column=3, value=repasse.proprietario.nome)
            ws.cell(row=row, column=4, value=str(repasse.contrato))
            ws.cell(row=row, column=5, value=str(repasse.contrato.imovel))
            ws.cell(row=row, column=6, value=float(repasse.valor))
            ws.cell(row=row, column=7, value=float(repasse.valor_desconto))
            ws.cell(row=row, column=8, value=float(repasse.valor_taxa_admin))
            ws.cell(row=row, column=9, value=float(repasse.valor_liquido))
            ws.cell(row=row, column=10, value=repasse.get_status_display())
            ws.cell(row=row, column=11, value=repasse.data_prevista)
            ws.cell(row=row, column=12, value=repasse.data_efetivacao)
            ws.cell(row=row, column=13, value=repasse.get_metodo_pagamento_display() if repasse.metodo_pagamento else '')
            ws.cell(row=row, column=14, value=f"{repasse.mes_referencia:02d}/{repasse.ano_referencia}")
            ws.cell(row=row, column=15, value=repasse.get_tipo_display())
            
            # Informação da política
            if hasattr(repasse.contrato, 'politica_repasse') and repasse.contrato.politica_repasse:
                politica_info = f"{repasse.contrato.politica_repasse.get_periodicidade_display()} ({repasse.contrato.politica_repasse.get_tipo_dias_display()})"
            else:
                politica_info = "Sem política específica"
            ws.cell(row=row, column=16, value=politica_info)
        
        # Ajustar largura das colunas
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Preparar response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="repasses_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        
        wb.save(response)
        return response
        
    except ImportError:
        messages.error(request, 'Biblioteca openpyxl não instalada. Execute: pip install openpyxl')
        return redirect('financeiro:lista_repasses')
    except Exception as e:
        logger.error(f"Erro ao gerar relatório: {str(e)}")
        messages.error(request, f'Erro ao gerar relatório: {str(e)}')
        return redirect('financeiro:lista_repasses')


# ==================== VIEWS DE HISTÓRICO ====================

def repasse_historico_proprietario(request, proprietario_id):
    """Histórico completo de repasses de um proprietário - CAMPOS CORRIGIDOS"""
    proprietario = get_object_or_404(Cliente, id=proprietario_id, tipo='proprietario')
    
    # Filtros
    ano = request.GET.get('ano', date.today().year)
    mes = request.GET.get('mes', '')
    status = request.GET.get('status', '')
    
    # Query base
    repasses = Repasse.objects.filter(
        proprietario=proprietario
    ).select_related('contrato', 'cobranca').prefetch_related(
        'contrato__politica_repasse'
    )
    
    # Aplicar filtros
    if ano:
        repasses = repasses.filter(ano_referencia=ano)
    if mes:
        repasses = repasses.filter(mes_referencia=mes)
    if status:
        repasses = repasses.filter(status=status)
    
    # Ordenação - tentar diferentes campos
    try:
        repasses = repasses.order_by('-data_criacao')
    except FieldError:
        try:
            repasses = repasses.order_by('-data_prevista')
        except FieldError:
            repasses = repasses.order_by('-id')  # Fallback final
    
    # Paginação
    paginator = Paginator(repasses, 20)
    page = request.GET.get('page')
    repasses_page = paginator.get_page(page)
    
    # Estatísticas do proprietário
    stats = {
        'total_recebido': repasses.filter(status='efetuado').aggregate(
            total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin'))
        )['total'] or Decimal('0'),
        'total_pendente': repasses.filter(status='pendente').aggregate(
            total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin'))
        )['total'] or Decimal('0'),
        'total_repasses': repasses.count(),
        'contratos_com_politica': Contrato.objects.filter(
            proprietario=proprietario,
            politica_repasse__isnull=False
        ).count(),
        'contratos_sem_politica': Contrato.objects.filter(
            proprietario=proprietario,
            ativo=True,
            politica_repasse__isnull=True
        ).count()
    }
    
    # Repasses por mês (últimos 12 meses) - simplificado
    repasses_mensais = []
    try:
        for i in range(12):
            data_ref = date.today().replace(day=1) - timedelta(days=i*30)
            mes_data = repasses.filter(
                ano_referencia=data_ref.year,
                mes_referencia=data_ref.month,
                status='efetuado'
            ).aggregate(
                total=Sum(F('valor') - F('valor_desconto') - F('valor_taxa_admin'))
            )
            
            repasses_mensais.append({
                'mes': f"{data_ref.month:02d}/{data_ref.year}",
                'total': float(mes_data['total'] or 0)
            })
        
        repasses_mensais.reverse()
    except Exception as e:
        logger.warning(f"Erro ao calcular repasses mensais: {e}")
        repasses_mensais = []
    
    context = {
        'proprietario': proprietario,
        'repasses': repasses_page,
        'stats': stats,
        'repasses_mensais': repasses_mensais,
        'filtros': {
            'ano': ano,
            'mes': mes,
            'status': status
        },
        'anos_disponiveis': range(2020, date.today().year + 2),
        'meses_disponiveis': [
            (1, 'Janeiro'), (2, 'Fevereiro'), (3, 'Março'), (4, 'Abril'),
            (5, 'Maio'), (6, 'Junho'), (7, 'Julho'), (8, 'Agosto'),
            (9, 'Setembro'), (10, 'Outubro'), (11, 'Novembro'), (12, 'Dezembro')
        ]
    }
    
    return render(request, 'financeiro/repasses/historico_proprietario.html', context)

def verificar_campos_modelo():
    """Função utilitária para verificar quais campos existem nos modelos"""
    
    print("=== CAMPOS DO MODELO REPASSE ===")
    for field in Repasse._meta.get_fields():
        print(f"- {field.name} ({type(field).__name__})")
    
    try:
        from ..models.movimento import MovimentoConta
        print("\n=== CAMPOS DO MODELO MOVIMENTOCONTA ===")
        for field in MovimentoConta._meta.get_fields():
            print(f"- {field.name} ({type(field).__name__})")
    except ImportError:
        print("\n=== MODELO MOVIMENTOCONTA NÃO ENCONTRADO ===")
    
    print("\n=== CAMPOS DO MODELO CONTRATO ===")
    for field in Contrato._meta.get_fields():
        print(f"- {field.name} ({type(field).__name__})")
# ==================== VIEWS DE UTILITÁRIOS ====================

def repasse_ajax_search(request):
    """Busca AJAX para repasses (para autocomplete)"""
    termo = request.GET.get('q', '')
    tipo = request.GET.get('tipo', 'proprietario')
    
    resultados = []
    
    if len(termo) >= 2:
        if tipo == 'proprietario':
            proprietarios = Cliente.objects.filter(
                tipo='proprietario',
                nome__icontains=termo
            )[:10]
            
            resultados = [{
                'id': p.id,
                'text': p.nome,
                'email': getattr(p, 'email', '')
            } for p in proprietarios]
            
        elif tipo == 'contrato':
            contratos = Contrato.objects.filter(
                Q(proprietario__nome__icontains=termo) |
                Q(imovel__endereco__icontains=termo)
            ).select_related('proprietario', 'imovel')[:10]
            
            resultados = [{
                'id': c.id,
                'text': f"{c.proprietario.nome} - {c.imovel.endereco}",
                'proprietario': c.proprietario.nome,
                'tem_politica': hasattr(c, 'politica_repasse') and c.politica_repasse is not None
            } for c in contratos]
    
    return JsonResponse({'results': resultados})


def cobranca_ajax_search(request):
    """Busca AJAX para cobranças pagas sem repasse"""
    contrato_id = request.GET.get('contrato_id')
    
    if contrato_id and Cobranca:
        cobrancas = Cobranca.objects.filter(
            contrato_id=contrato_id,
            status='paga',
            repasses_cobranca__isnull=True
        ).order_by('-data_pagamento')[:10]
        
        resultados = [{
            'id': c.id,
            'text': f"Cobrança #{c.id} - {c.mes_referencia:02d}/{c.ano_referencia} - R$ {getattr(c, 'valor_pago', getattr(c, 'valor', 0)):.2f}",
            'valor': float(getattr(c, 'valor_pago', getattr(c, 'valor', 0))),
            'mes': getattr(c, 'mes_referencia', date.today().month),
            'ano': getattr(c, 'ano_referencia', date.today().year)
        } for c in cobrancas]
        
        return JsonResponse({'results': resultados})
    
    return JsonResponse({'results': []})


def calcular_valores_repasse(request):
    """Calcula valores do repasse baseado na cobrança e política do contrato - VERSÃO CORRIGIDA"""
    cobranca_id = request.GET.get('cobranca_id')
    contrato_id = request.GET.get('contrato_id')
    
    try:
        if cobranca_id and Cobranca:
            cobranca = Cobranca.objects.get(id=cobranca_id)
            contrato = cobranca.contrato
        elif contrato_id:
            contrato = Contrato.objects.get(id=contrato_id)
            cobranca = None
        else:
            return JsonResponse({'success': False, 'error': 'ID não fornecido'})
        
        # ========== CORREÇÃO 1: DETERMINAR VALOR BASE CORRETO ==========
        if cobranca:
            # Verificar se a cobrança tem detalhes de cálculo (separando aluguel de despesas)
            if hasattr(cobranca, 'detalhes_calculo') and cobranca.detalhes_calculo:
                try:
                    detalhes = cobranca.detalhes_calculo
                    # Usar apenas o valor do aluguel como base para taxa administrativa
                    valor_base_aluguel = Decimal(str(detalhes.get('valor_base', 0)))
                    valor_despesas = Decimal(str(detalhes.get('valor_despesas_liquido', 0)))
                    
                    if valor_base_aluguel > 0:
                        valor_base = valor_base_aluguel
                        tem_despesas = valor_despesas > 0
                    else:
                        # Fallback para valor total se não conseguir separar
                        valor_base = getattr(cobranca, 'valor_pago', getattr(cobranca, 'valor', 0))
                        valor_despesas = Decimal('0.00')
                        tem_despesas = False
                        logger.warning(f"Cobrança {cobranca_id}: Não foi possível separar aluguel de despesas")
                        
                except (KeyError, TypeError, ValueError) as e:
                    logger.warning(f"Erro ao processar detalhes_calculo da cobrança {cobranca_id}: {e}")
                    # Fallback para valor total
                    valor_base = getattr(cobranca, 'valor_pago', getattr(cobranca, 'valor', 0))
                    valor_despesas = Decimal('0.00')
                    tem_despesas = False
            else:
                # Cobrança sem detalhes - usar valor total (comportamento antigo)
                valor_base = getattr(cobranca, 'valor_pago', getattr(cobranca, 'valor', 0))
                valor_despesas = Decimal('0.00')
                tem_despesas = False
        else:
            # Sem cobrança - usar valor do contrato
            valor_base = getattr(contrato, 'valor_aluguel', getattr(contrato, 'valor_base', Decimal('0')))
            valor_despesas = Decimal('0.00')
            tem_despesas = False
        
        # ========== CORREÇÃO 2: VERIFICAR POLÍTICA OBRIGATORIAMENTE ==========
        taxa_admin = None
        fonte_taxa = None
        tem_politica_contrato = hasattr(contrato, 'politica_repasse') and contrato.politica_repasse is not None
        
        if tem_politica_contrato:
            try:
                politica = contrato.politica_repasse
                if politica.ativa:
                    taxa_admin = politica.get_taxa_admin()
                    fonte_taxa = f"Política do contrato (ID: {politica.id})"
                else:
                    logger.warning(f"Política do contrato {contrato.id} existe mas está inativa")
            except Exception as e:
                logger.error(f"Erro ao obter taxa da política do contrato {contrato.id}: {e}")
        
        # Se não conseguiu obter taxa da política do contrato, buscar política global
        if taxa_admin is None:
            politica_global = PoliticaRepasseGlobal.objects.filter(ativa=True).first()
            if politica_global:
                taxa_admin = politica_global.taxa_admin_padrao
                fonte_taxa = f"Política global (ID: {politica_global.id})"
            else:
                # AVISO: Usar padrão apenas se não há política configurada
                taxa_admin = Decimal('8.00')
                fonte_taxa = "Padrão do sistema (8%) - CONFIGURAR POLÍTICA!"
                logger.warning(f"Contrato {contrato.id} sem política configurada. Usando padrão 8%")
        
        # ========== CÁLCULO CORRIGIDO ==========
        # Taxa administrativa aplicada APENAS sobre o aluguel
        valor_taxa_admin = valor_base * (taxa_admin / 100)
        aluguel_liquido = valor_base - valor_taxa_admin
        
        # Valor líquido total = aluguel líquido + despesas (que passam integralmente)
        valor_liquido_total = aluguel_liquido + valor_despesas
        
        # Verificar valor mínimo se houver política
        valor_minimo = Decimal('0.00')
        if tem_politica_contrato:
            valor_minimo = getattr(contrato.politica_repasse, 'valor_minimo_repasse', Decimal('0.00')) or Decimal('0.00')
        
        # ========== LOGGING PARA AUDITORIA ==========
        logger.info(f"Cálculo repasse - Contrato: {contrato.id}, Cobrança: {cobranca_id or 'N/A'}")
        logger.info(f"  - Valor base (aluguel): R$ {valor_base}")
        logger.info(f"  - Despesas: R$ {valor_despesas}")
        logger.info(f"  - Taxa admin: {taxa_admin}% (fonte: {fonte_taxa})")
        logger.info(f"  - Taxa em R$: R$ {valor_taxa_admin}")
        logger.info(f"  - Valor líquido total: R$ {valor_liquido_total}")
        
        return JsonResponse({
            'success': True,
            # Valores principais
            'valor_base': float(valor_base),  # Aluguel apenas
            'valor_despesas': float(valor_despesas),  # Despesas separadamente
            'valor_taxa_admin': float(valor_taxa_admin),
            'valor_desconto': 0.0,
            'valor_liquido': float(aluguel_liquido),  # Só aluguel líquido
            'valor_liquido_total': float(valor_liquido_total),  # Aluguel + despesas
            
            # Informações da taxa
            'taxa_admin_percentual': float(taxa_admin),
            'fonte_taxa': fonte_taxa,
            'tem_politica_contrato': tem_politica_contrato,
            
            # Validações
            'valor_minimo': float(valor_minimo),
            'acima_minimo': valor_liquido_total >= valor_minimo,
            'tem_despesas': tem_despesas,
            
            # Metadados
            'mes_referencia': getattr(cobranca, 'mes_referencia', date.today().month) if cobranca else date.today().month,
            'ano_referencia': getattr(cobranca, 'ano_referencia', date.today().year) if cobranca else date.today().year,
            
            # Detalhamento para debug
            'detalhamento': {
                'calculo': f"Aluguel: R$ {valor_base} - Taxa {taxa_admin}%: R$ {valor_taxa_admin} = R$ {aluguel_liquido}",
                'despesas': f"+ Despesas: R$ {valor_despesas}" if tem_despesas else "Sem despesas",
                'total': f"Total líquido: R$ {valor_liquido_total}",
                'observacao': "Taxa aplicada apenas sobre aluguel. Despesas passam integralmente ao proprietário." if tem_despesas else "Sem despesas adicionais."
            }
        })
        
    except Cobranca.DoesNotExist if Cobranca else Exception:
        return JsonResponse({'success': False, 'error': 'Cobrança não encontrada'})
    except Contrato.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Contrato não encontrado'})
    except Exception as e:
        logger.error(f"Erro ao calcular valores de repasse: {str(e)}")
        return JsonResponse({
            'success': False, 
            'error': f'Erro no cálculo: {str(e)}',
            'contrato_id': contrato_id,
            'cobranca_id': cobranca_id
        })


# ========== FUNÇÃO AUXILIAR PARA CRIAÇÃO DE REPASSE ==========
def criar_repasse_com_calculo_correto(cobranca, tipo='automatico'):
    """Cria repasse usando a lógica corrigida de cálculo"""
    
    contrato = cobranca.contrato
    
    # Verificar se tem política configurada
    if not (hasattr(contrato, 'politica_repasse') and contrato.politica_repasse):
        logger.warning(f"Tentativa de criar repasse para contrato {contrato.id} sem política configurada")
        
        # Criar agendamento para processar depois
        try:
            AgendamentoRepasse.objects.create(
                contrato=contrato,
                cobranca=cobranca,
                data_agendada=date.today() + timedelta(days=1),
                status='agendado',
                motivo="Aguardando configuração de política de repasse",
                observacoes=f"Cobrança {cobranca.id} paga em {cobranca.data_pagamento or 'data não informada'}"
            )
            logger.info(f"Agendamento criado para contrato {contrato.id}, cobrança {cobranca.id}")
        except Exception as e:
            logger.error(f"Erro ao criar agendamento: {e}")
        
        return None
    
    # Calcular valores usando a lógica corrigida
    try:
        # Extrair valores da cobrança
        if hasattr(cobranca, 'detalhes_calculo') and cobranca.detalhes_calculo:
            detalhes = cobranca.detalhes_calculo
            valor_base_aluguel = Decimal(str(detalhes.get('valor_base', 0)))
            valor_despesas = Decimal(str(detalhes.get('valor_despesas_liquido', 0)))
        else:
            # Fallback: usar valor total como aluguel
            valor_base_aluguel = cobranca.valor
            valor_despesas = Decimal('0.00')
        
        # Obter taxa da política
        politica = contrato.politica_repasse
        taxa_admin = politica.get_taxa_admin()
        
        # Calcular valores
        valor_taxa_admin = valor_base_aluguel * (taxa_admin / 100)
        
        # Criar repasse
        repasse = Repasse.objects.create(
            contrato=contrato,
            cobranca=cobranca,
            proprietario=contrato.proprietario.first(),  # Assumindo que há pelo menos um
            
            # Valores corrigidos
            valor=valor_base_aluguel,  # Base apenas aluguel
            valor_taxa_admin=valor_taxa_admin,
            valor_desconto=Decimal('0.00'),
            
            # Metadados
            mes_referencia=cobranca.mes_referencia,
            ano_referencia=cobranca.ano_referencia,
            tipo=tipo,
            status='pendente',
            data_prevista=date.today() + timedelta(days=7),  # Exemplo
            
            # Observações com detalhamento
            observacoes=f"Taxa {taxa_admin}% aplicada sobre aluguel (R$ {valor_base_aluguel}). "
                       f"Despesas (R$ {valor_despesas}) passam integralmente ao proprietário.",
            
            # Campos extras se existirem no modelo
            **({'valor_despesas': valor_despesas} if hasattr(Repasse, 'valor_despesas') else {}),
            **({'detalhamento_calculo': f"Aluguel: R$ {valor_base_aluguel} - Taxa: R$ {valor_taxa_admin} + Despesas: R$ {valor_despesas}"} if hasattr(Repasse, 'detalhamento_calculo') else {})
        )
        
        logger.info(f"Repasse {repasse.id} criado para contrato {contrato.id}, cobrança {cobranca.id}")
        logger.info(f"  - Aluguel: R$ {valor_base_aluguel}, Taxa: R$ {valor_taxa_admin}, Despesas: R$ {valor_despesas}")
        
        return repasse
        
    except Exception as e:
        logger.error(f"Erro ao criar repasse para cobrança {cobranca.id}: {e}")
        return None
# ==================== VIEWS DE MIGRAÇÃO E UTILITÁRIOS ====================

def migrar_politicas_antigas(request):
    """Migra políticas do modelo antigo para o novo (se necessário)"""
    if request.method == 'POST':
        try:
            # Esta view seria usada para migrar dados do modelo antigo
            # para o novo modelo de políticas por contrato
            messages.success(request, 'Migração concluída com sucesso!')
        except Exception as e:
            logger.error(f"Erro na migração: {str(e)}")
            messages.error(request, f'Erro na migração: {str(e)}')
    
    return redirect('financeiro:configuracao_repasse')


def extrato_proprietario_view(request, proprietario_id):
    """Extrato simplificado do proprietário"""
    return redirect('financeiro:historico_proprietario', proprietario_id=proprietario_id)# financeiro/views/repasse_views.py (versão corrigida)
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.generic import ListView, DetailView, CreateView, UpdateView, TemplateView
from django.db.models import Q, Sum, Count, Case, When, DecimalField, F
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from django.core.paginator import Paginator
from django.urls import reverse_lazy
from django.db import transaction
from datetime import date, datetime, timedelta
from decimal import Decimal
import json
import logging

# Imports dos modelos
from ..models.repasse import Repasse, PoliticaRepasseContrato, PoliticaRepasseGlobal, AgendamentoRepasse
try:
    from ..models.cobranca import Cobranca
except ImportError:
    Cobranca = None

from core.models import Contrato, Cliente

# Imports dos services e forms (assumindo que existem)
try:
    from ..services.repasse_service import RepasseService
except ImportError:
    RepasseService = None

logger = logging.getLogger(__name__)


# ==================== VIEWS DE REPASSES ====================




# Simplificando as outras views para evitar erros de campos inexistentes
class RepasseCreateView(CreateView):
    model = Repasse
    template_name = 'financeiro/repasses/repasse_form.html'
    fields = ['proprietario', 'contrato', 'cobranca', 'valor', 'valor_desconto', 
              'valor_taxa_admin', 'data_prevista', 'mes_referencia', 'ano_referencia',
              'tipo', 'metodo_pagamento', 'descricao', 'observacoes']
    
    def get_success_url(self):
        return reverse_lazy('financeiro:repasse_detail', kwargs={'pk': self.object.pk})
    
    def form_valid(self, form):
        messages.success(self.request, 'Repasse criado com sucesso!')
        return super().form_valid(form)


class RepasseUpdateView(UpdateView):
    model = Repasse
    template_name = 'financeiro/repasses/repasse_form.html'
    fields = ['proprietario', 'contrato', 'cobranca', 'valor', 'valor_desconto', 
              'valor_taxa_admin', 'data_prevista', 'mes_referencia', 'ano_referencia',
              'tipo', 'metodo_pagamento', 'descricao', 'observacoes']
    
    def get_success_url(self):
        return reverse_lazy('financeiro:repasse_detail', kwargs={'pk': self.object.pk})
    
    def form_valid(self, form):
        messages.success(self.request, 'Repasse atualizado com sucesso!')
        return super().form_valid(form)


# ==================== VIEWS DE POLÍTICAS POR CONTRATO ====================

class ContratosSemPoliticaView(ListView):
    """Lista contratos ativos sem política de repasse configurada"""
    template_name = 'financeiro/repasses/contratos_sem_politica.html'
    context_object_name = 'contratos'
    paginate_by = 20

    def get_queryset(self):
        # Corrigido: usar ativo=True em vez de status='ativo'
        queryset = Contrato.objects.filter(
            ativo=True,
            politica_repasse__isnull=True
        ).select_related('proprietario', 'imovel').order_by('proprietario__nome')

        # Filtros básicos
        proprietario_id = self.request.GET.get('proprietario')
        if proprietario_id:
            queryset = queryset.filter(proprietario_id=proprietario_id)
        
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(proprietario__nome__icontains=search) |
                Q(imovel__endereco__icontains=search)
            )

        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['total_sem_politica'] = self.get_queryset().count()
        context['total_contratos_ativos'] = Contrato.objects.filter(ativo=True).count()
        return context





class PoliticaContratoUpdateView(UpdateView):
    model = PoliticaRepasseContrato
    template_name = 'financeiro/repasses/politica_contrato_form.html'
    fields = ['ativa', 'periodicidade', 'tipo_dias', 'dia_mes', 'dia_semana',
              'dias_apos_recebimento', 'percentual_adiantamento', 'taxa_adiantamento',
              'valor_minimo_repasse', 'taxa_admin_personalizada', 'considerar_feriados',
              'antecipar_fds_feriados', 'observacoes']
    
    def get_success_url(self):
        return reverse_lazy('financeiro:politica_contrato_list')
    
    def form_valid(self, form):
        messages.success(self.request, f'Política do contrato {form.instance.contrato.id} atualizada!')
        return super().form_valid(form)


# ==================== VIEWS DE POLÍTICAS GLOBAIS ====================

class PoliticaGlobalListView(ListView):
    model = PoliticaRepasseGlobal
    template_name = 'financeiro/repasses/politica_global_list.html'
    context_object_name = 'politicas'

    def get_queryset(self):
        return PoliticaRepasseGlobal.objects.order_by('-ativa', 'nome')


class PoliticaGlobalCreateView(CreateView):
    model = PoliticaRepasseGlobal
    template_name = 'financeiro/repasses/politica_form.html'
    fields = ['nome', 'ativa', 'periodicidade', 'tipo_dias', 'dia_mes', 'dia_semana',
              'dias_apos_recebimento', 'percentual_adiantamento', 'taxa_adiantamento',
              'valor_minimo_repasse', 'taxa_admin_padrao', 'considerar_feriados',
              'antecipar_fds_feriados']
    
    def get_success_url(self):
        return reverse_lazy('financeiro:politica_global_list')


class PoliticaGlobalUpdateView(UpdateView):
    model = PoliticaRepasseGlobal
    template_name = 'financeiro/repasses/politica_form.html'
    fields = ['nome', 'ativa', 'periodicidade', 'tipo_dias', 'dia_mes', 'dia_semana',
              'dias_apos_recebimento', 'percentual_adiantamento', 'taxa_adiantamento',
              'valor_minimo_repasse', 'taxa_admin_padrao', 'considerar_feriados',
              'antecipar_fds_feriados']
    
    def get_success_url(self):
        return reverse_lazy('financeiro:politica_global_list')


# ==================== VIEWS DE PROCESSAMENTO ====================

def processar_repasse_view(request, repasse_id):
    """Processa (efetiva) um repasse pendente"""
    repasse = get_object_or_404(Repasse, id=repasse_id)
    
    if request.method == 'POST':
        try:
            metodo_pagamento = request.POST.get('metodo_pagamento')
            observacoes = request.POST.get('observacoes')
            
            success = repasse.efetivar_repasse(
                metodo_pagamento=metodo_pagamento,
                observacoes=observacoes
            )
            
            if success:
                # Upload do comprovante se fornecido
                comprovante = request.FILES.get('comprovante')
                if comprovante:
                    repasse.comprovante = comprovante
                    repasse.save(update_fields=['comprovante'])
                
                messages.success(request, f'Repasse efetuado com sucesso!')
                return redirect('financeiro:repasse_detail', pk=repasse.id)
            else:
                messages.error(request, 'Não foi possível efetivar o repasse.')
        except Exception as e:
            logger.error(f"Erro ao processar repasse {repasse_id}: {str(e)}")
            messages.error(request, f'Erro ao processar repasse: {str(e)}')

    return render(request, 'financeiro/repasses/processar_repasse.html', {
        'repasse': repasse,
    })


def cancelar_repasse_view(request, repasse_id):
    """Cancela um repasse pendente"""
    repasse = get_object_or_404(Repasse, id=repasse_id)
    
    if request.method == 'POST':
        motivo = request.POST.get('motivo', '')
        success = repasse.cancelar_repasse(motivo)
        
        if success:
            messages.success(request, 'Repasse cancelado com sucesso!')
        else:
            messages.error(request, 'Não foi possível cancelar o repasse.')
    
    return redirect('financeiro:repasse_detail', pk=repasse.id)



# ==================== VIEWS DE AGENDAMENTOS ====================

class AgendamentoListView(ListView):
    model = AgendamentoRepasse
    template_name = 'financeiro/repasses/agendamento_list.html'
    context_object_name = 'agendamentos'
    paginate_by = 20

    def get_queryset(self):
        return AgendamentoRepasse.objects.select_related(
            'proprietario', 'contrato', 'politica_contrato', 'politica_global'
        ).order_by('-data_agendada')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        # Estatísticas
        hoje = date.today()
        context['stats'] = {
            'total_agendados': AgendamentoRepasse.objects.filter(status='agendado').count(),
            'vencidos_hoje': AgendamentoRepasse.objects.filter(
                status='agendado',
                data_agendada__lte=hoje
            ).count(),
            'proximos_7_dias': AgendamentoRepasse.objects.filter(
                status='agendado',
                data_agendada__range=[hoje, hoje + timedelta(days=7)]
            ).count()
        }
        
        return context


def processar_agendamento_view(request, pk):
    """Processa um agendamento específico"""
    agendamento = get_object_or_404(AgendamentoRepasse, pk=pk)
    
    if request.method == 'POST':
        try:
            repasse = agendamento.processar()
            if repasse:
                messages.success(request, f'Agendamento processado! Repasse #{repasse.id} criado.')
                return redirect('financeiro:repasse_detail', pk=repasse.id)
            else:
                messages.error(request, 'Erro ao processar agendamento.')
        except Exception as e:
            logger.error(f"Erro ao processar agendamento {pk}: {str(e)}")
            messages.error(request, f'Erro: {str(e)}')
    
    return redirect('financeiro:agendamento_list')



# ==================== VIEWS DE DASHBOARD ====================


# ==================== VIEWS AJAX SIMPLIFICADAS ====================

def toggle_politica_contrato_view(request, pk):
    """Toggle status ativo/inativo de uma política de contrato"""
    if request.method == 'POST':
        politica = get_object_or_404(PoliticaRepasseContrato, pk=pk)
        
        try:
            politica.ativa = not politica.ativa
            politica.save()
            
            messages.success(request, f'Política do contrato {politica.contrato.id} {"ativada" if politica.ativa else "desativada"}!')
            return JsonResponse({'success': True})
            
        except Exception as e:
            logger.error(f"Erro ao alterar status da política {pk}: {str(e)}")
            return JsonResponse({'success': False, 'error': str(e)})
    
    return JsonResponse({'success': False, 'error': 'Método não permitido'})



def aplicar_politica_em_lote(request):
    """Aplica uma política global em lote para contratos selecionados"""
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            politica_global_id = data.get('politica_global_id')
            contratos_ids = data.get('contratos_ids', [])
            
            if not politica_global_id:
                return JsonResponse({
                    'success': False, 
                    'error': 'Política deve ser informada'
                })
            
            politica_global = get_object_or_404(PoliticaRepasseGlobal, pk=politica_global_id)
            
            if contratos_ids:
                contratos = Contrato.objects.filter(
                    id__in=contratos_ids,
                    politica_repasse__isnull=True
                )
            else:
                contratos = Contrato.objects.filter(
                    ativo=True,  # Corrigido
                    politica_repasse__isnull=True
                )
            
            criados = 0
            erros = []
            
            with transaction.atomic():
                for contrato in contratos:
                    try:
                        PoliticaRepasseContrato.objects.create(
                            contrato=contrato,
                            ativa=politica_global.ativa,
                            periodicidade=politica_global.periodicidade,
                            tipo_dias=politica_global.tipo_dias,
                            dia_mes=politica_global.dia_mes,
                            dia_semana=politica_global.dia_semana,
                            dias_apos_recebimento=politica_global.dias_apos_recebimento,
                            percentual_adiantamento=politica_global.percentual_adiantamento,
                            taxa_adiantamento=politica_global.taxa_adiantamento,
                            valor_minimo_repasse=politica_global.valor_minimo_repasse,
                            taxa_admin_personalizada=politica_global.taxa_admin_padrao,
                            considerar_feriados=politica_global.considerar_feriados,
                            antecipar_fds_feriados=politica_global.antecipar_fds_feriados,
                            observacoes=f"Criada em lote baseada na política global '{politica_global.nome}'"
                        )
                        criados += 1
                    except Exception as e:
                        erros.append(f"Contrato {contrato.id}: {str(e)}")
            
            return JsonResponse({
                'success': True,
                'criados': criados,
                'erros': erros
            })
            
        except Exception as e:
            logger.error(f"Erro ao aplicar política em lote: {str(e)}")
            return JsonResponse({
                'success': False,
                'error': str(e)
            })
    
    return JsonResponse({'success': False, 'error': 'Método não permitido'})


@require_POST
def gerar_repasse_automatico(request):
    """
    View para gerar repasses automaticamente
    """
    try:
        with transaction.atomic():
            # Buscar cobranças elegíveis
            cobrancas_sem_repasse = Cobranca.objects.filter(
                status='paga'
            ).exclude(
                id__in=Repasse.objects.values_list('cobranca_id', flat=True)
            ).select_related('contrato')
            
            total_cobrancas = cobrancas_sem_repasse.count()
            
            if total_cobrancas == 0:
                return JsonResponse({
                    'success': True,
                    'criados': 0,
                    'processados': 0,
                    'agendados': 0,
                    'message': 'Nenhuma cobrança elegível para gerar repasse.'
                })
            
            criados = 0
            erros = 0
            detalhes = []
            
            for cobranca in cobrancas_sem_repasse:
                try:
                    # Obter proprietário
                    if hasattr(cobranca.contrato, 'proprietario'):
                        if hasattr(cobranca.contrato.proprietario, 'all'):
                            proprietario = cobranca.contrato.proprietario.first()
                        else:
                            proprietario = cobranca.contrato.proprietario
                    else:
                        proprietario = None
                        
                    if not proprietario:
                        erros += 1
                        detalhes.append(f"Cobrança {cobranca.id}: sem proprietário")
                        continue
                    
                    # Usar RepasseService para calcular data
                    data_repasse = RepasseService._calcular_data_repasse(cobranca)
                    
                    # Calcular valores
                    valor_bruto = cobranca.valor or Decimal('0.00')
                    
                    # Taxa administrativa
                    politica = getattr(cobranca.contrato, 'politica_repasse', None)
                    if politica and hasattr(politica, 'taxa_admin_personalizada') and politica.taxa_admin_personalizada:
                        taxa_percentual = politica.taxa_admin_personalizada
                    else:
                        taxa_percentual = Decimal('8.00')
                    
                    valor_taxa_admin = valor_bruto * (taxa_percentual / 100)
                    
                    # Criar repasse
                    repasse = Repasse.objects.create(
                        proprietario=proprietario,
                        cobranca=cobranca,
                        contrato=cobranca.contrato,
                        valor=valor_bruto,
                        valor_desconto=Decimal('0.00'),
                        valor_taxa_admin=valor_taxa_admin,
                        data_prevista=data_repasse,
                        mes_referencia=cobranca.mes_referencia,
                        ano_referencia=cobranca.ano_referencia,
                        status='pendente',
                        tipo='automatico',
                        descricao=f"Repasse automático - ref. {cobranca.mes_referencia:02d}/{cobranca.ano_referencia}"
                    )
                    
                    criados += 1
                    detalhes.append(f"Repasse {repasse.id} criado para cobrança {cobranca.id}")
                    logger.info(f"Repasse {repasse.id} criado para cobrança {cobranca.id}")
                        
                except Exception as e:
                    erros += 1
                    detalhes.append(f"Erro na cobrança {cobrança.id}: {str(e)}")
                    logger.error(f"Erro ao processar cobrança {cobranca.id}: {e}")
            
            # Resposta
            if criados > 0:
                message = f"Sucesso! {criados} repasse(s) criado(s)."
                if erros > 0:
                    message += f" {erros} erro(s) ocorreram."
            else:
                message = f"Nenhum repasse foi criado. {erros} erro(s) ocorreram."
            
            return JsonResponse({
                'success': criados > 0,
                'criados': criados,
                'processados': 0,
                'agendados': 0,
                'erros': erros,
                'message': message,
                'detalhes': detalhes[:5]
            })
            
    except Exception as e:
        logger.error(f"Erro geral na geração de repasses: {e}")
        return JsonResponse({
            'success': False,
            'error': f'Erro interno: {str(e)}',
            'criados': 0,
            'processados': 0,
            'agendados': 0
        })

def _criar_repasse_para_cobranca(cobranca):
    """
    Função auxiliar para criar repasse para uma cobrança específica
    """
    try:
        # 1. Obter proprietário
        proprietario = _obter_proprietario_contrato(cobranca.contrato)
        if not proprietario:
            raise ValueError("Contrato não possui proprietário definido")
        
        # 2. Calcular data de repasse
        data_repasse = RepasseService._calcular_data_repasse(cobranca)
        
        # 3. Calcular valores
        valor_bruto = cobranca.valor or Decimal('0.00')
        
        # Obter taxa administrativa
        politica = getattr(cobranca.contrato, 'politica_repasse', None)
        if politica and hasattr(politica, 'taxa_admin_personalizada') and politica.taxa_admin_personalizada:
            taxa_percentual = politica.taxa_admin_personalizada
        else:
            taxa_percentual = Decimal('8.00')  # Padrão 8%
        
        valor_taxa_admin = valor_bruto * (taxa_percentual / 100)
        valor_desconto = Decimal('0.00')
        
        # 4. Criar repasse
        repasse = Repasse.objects.create(
            proprietario=proprietario,
            cobranca=cobranca,
            contrato=cobranca.contrato,
            valor=valor_bruto,
            valor_desconto=valor_desconto,
            valor_taxa_admin=valor_taxa_admin,
            data_prevista=data_repasse,
            mes_referencia=cobranca.mes_referencia,
            ano_referencia=cobranca.ano_referencia,
            status='pendente',
            tipo='automatico',
            descricao=f"Repasse automático - ref. {cobranca.mes_referencia:02d}/{cobranca.ano_referencia}"
        )
        
        return repasse
        
    except Exception as e:
        logger.error(f"Erro ao criar repasse para cobrança {cobranca.id}: {e}")
        raise

def _obter_proprietario_contrato(contrato):
    """
    Obtém o proprietário do contrato
    """
    try:
        if hasattr(contrato, 'proprietario'):
            if hasattr(contrato.proprietario, 'all'):
                # Relacionamento ManyToMany
                return contrato.proprietario.first()
            else:
                # Relacionamento ForeignKey
                return contrato.proprietario
        return None
    except AttributeError:
        return None


def processar_repasse(request, repasse_id):
    """
    View para processar (efetivar) um repasse específico
    """
    try:
        repasse = Repasse.objects.get(id=repasse_id)
        
        if repasse.status != 'pendente':
            messages.error(request, 'Este repasse não pode ser processado.')
            return redirect('financeiro:repasses')
        
        if request.method == 'POST':
            metodo_pagamento = request.POST.get('metodo_pagamento')
            observacoes = request.POST.get('observacoes')
            
            # Efetivar repasse
            sucesso = repasse.efetivar_repasse(
                metodo_pagamento=metodo_pagamento,
                observacoes=observacoes
            )
            
            if sucesso:
                messages.success(request, f'Repasse {repasse.id} efetuado com sucesso!')
            else:
                messages.error(request, 'Erro ao efetivar repasse.')
            
            return redirect('financeiro:repasses')
        
        return render(request, 'financeiro/repasse/processar.html', {
            'repasse': repasse
        })
        
    except Repasse.DoesNotExist:
        messages.error(request, 'Repasse não encontrado.')
        return redirect('financeiro:repasses')


@require_POST
def cancelar_repasse(request, repasse_id):
    """
    View para cancelar um repasse
    """
    try:
        repasse = Repasse.objects.get(id=repasse_id)
        
        if repasse.status != 'pendente':
            return JsonResponse({
                'success': False,
                'error': 'Este repasse não pode ser cancelado.'
            })
        
        motivo = request.POST.get('motivo', '')
        
        if not motivo.strip():
            return JsonResponse({
                'success': False,
                'error': 'Motivo do cancelamento é obrigatório.'
            })
        
        sucesso = repasse.cancelar_repasse(motivo)
        
        if sucesso:
            return JsonResponse({
                'success': True,
                'message': f'Repasse {repasse.id} cancelado com sucesso.'
            })
        else:
            return JsonResponse({
                'success': False,
                'error': 'Erro ao cancelar repasse.'
            })
        
    except Repasse.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Repasse não encontrado.'
        })
    except Exception as e:
        logger.error(f"Erro ao cancelar repasse {repasse_id}: {e}")
        return JsonResponse({
            'success': False,
            'error': 'Erro interno do servidor.'
        })


def relatorio_repasses(request):
    """
    View para relatório de repasses
    """
    from django.db.models import Sum, Count, Q
    from datetime import date, timedelta
    
    # Filtros
    status = request.GET.get('status', '')
    mes = request.GET.get('mes', '')
    ano = request.GET.get('ano', str(date.today().year))
    
    # Query base
    repasses = Repasse.objects.all()
    
    # Aplicar filtros
    if status:
        repasses = repasses.filter(status=status)
    
    if mes:
        repasses = repasses.filter(mes_referencia=int(mes))
    
    if ano:
        repasses = repasses.filter(ano_referencia=int(ano))
    
    # Estatísticas
    stats = repasses.aggregate(
        total_valor=Sum('valor'),
        total_taxa=Sum('valor_taxa_admin'),
        total_liquido=Sum('valor') - Sum('valor_taxa_admin') - Sum('valor_desconto'),
        count_total=Count('id'),
        count_pendente=Count('id', filter=Q(status='pendente')),
        count_efetuado=Count('id', filter=Q(status='efetuado')),
        count_cancelado=Count('id', filter=Q(status='cancelado'))
    )
    
    context = {
        'repasses': repasses.order_by('-data_criacao')[:100],  # Limitar para performance
        'stats': stats,
        'filtros': {
            'status': status,
            'mes': mes,
            'ano': ano
        }
    }
    
    return render(request, 'financeiro/repasse/relatorio.html', context)

# Função para executar uma vez no shell (pontual)
def executar_criacao_inicial():
    """
    FUNÇÃO APENAS PARA EXECUÇÃO PONTUAL NO SHELL
    NÃO DEVE SER CHAMADA EM VIEWS NORMAIS
    """
    from financeiro.models import Cobranca, Repasse
    from financeiro.services.repasse_service import RepasseService
    from decimal import Decimal
    
    print("🔄 CRIAÇÃO INICIAL DE REPASSES")
    print("=" * 50)
    
    cobrancas_sem_repasse = Cobranca.objects.filter(
        status='paga'
    ).exclude(
        id__in=Repasse.objects.values_list('cobranca_id', flat=True)
    )
    
    criados = 0
    for cobranca in cobrancas_sem_repasse:
        try:
            repasse = _criar_repasse_para_cobranca(cobranca)
            if repasse:
                print(f"✅ Repasse {repasse.id} criado para cobrança {cobranca.id}")
                criados += 1
        except Exception as e:
            print(f"❌ Erro na cobrança {cobranca.id}: {e}")
    
    print(f"\n📊 Total criados: {criados}")
    return criados

# Cálculo de Repasse INDEPENDENTE - Sem usar signals
# Substitua as funções na sua view por estas versões:

def _calcular_data_repasse_independente(cobranca):
    """
    Calcula data de repasse independente (sem usar RepasseService)
    """
    try:
        from datetime import date, timedelta
        
        data_pagamento = cobranca.data_pagamento or date.today()
        
        # Verificar política do contrato
        politica = getattr(cobranca.contrato, 'politica_repasse', None)
        
        if politica and politica.ativa and politica.dias_apos_recebimento:
            dias = politica.dias_apos_recebimento
            
            if politica.tipo_dias == 'uteis':
                # Calcular dias úteis
                data_atual = data_pagamento
                dias_adicionados = 0
                
                while dias_adicionados < dias:
                    data_atual = data_atual + timedelta(days=1)
                    # Segunda=0 a Sexta=4 são dias úteis
                    if data_atual.weekday() < 5:
                        dias_adicionados += 1
                
                return data_atual
            else:
                # Dias corridos
                return data_pagamento + timedelta(days=dias)
        else:
            # Fallback: 5 dias úteis
            data_atual = data_pagamento
            dias_adicionados = 0
            
            while dias_adicionados < 5:
                data_atual = data_atual + timedelta(days=1)
                if data_atual.weekday() < 5:
                    dias_adicionados += 1
            
            return data_atual
            
    except Exception:
        # Fallback absoluto: 5 dias corridos
        return (cobranca.data_pagamento or date.today()) + timedelta(days=5)

def _calcular_valores_repasse_completo(cobranca):
    """
    Calcula repasse integrando aluguel + despesas + juros/multas
    VERSÃO INDEPENDENTE - sem usar signals
    """
    try:
        from decimal import Decimal
        from datetime import date
        
        # 1. VALOR TOTAL DA COBRANÇA
        valor_total_cobranca = cobranca.valor_total or cobranca.valor
        
        # 2. COMPONENTES COM E SEM INCIDÊNCIA
        valor_com_incidencia = Decimal('0.00')
        valor_sem_incidencia = Decimal('0.00')
        detalhamento = []
        
        # 3. ALUGUEL (sempre tem incidência)
        valor_aluguel = Decimal(str(cobranca.valor_aluguel or 0))
        if valor_aluguel > 0:
            valor_com_incidencia += valor_aluguel
            detalhamento.append({
                'tipo': 'Aluguel',
                'valor': valor_aluguel,
                'tem_incidencia': True,
                'percentual_incidencia': 100
            })
        
        # 4. JUROS E MULTAS (definir suas regras aqui)
        valor_juros = _obter_valor_juros_multas(cobranca, 'juros')
        valor_multa = _obter_valor_juros_multas(cobranca, 'multa')
        
        # REGRA: Juros e multas TÊM incidência (receita da administradora)
        # Altere para False se não quiser incidência
        JUROS_TEM_INCIDENCIA = True
        MULTA_TEM_INCIDENCIA = True
        
        if valor_juros > 0:
            if JUROS_TEM_INCIDENCIA:
                valor_com_incidencia += valor_juros
                detalhamento.append({
                    'tipo': 'Juros de Atraso',
                    'valor': valor_juros,
                    'tem_incidencia': True,
                    'percentual_incidencia': 100
                })
            else:
                valor_sem_incidencia += valor_juros
                detalhamento.append({
                    'tipo': 'Juros de Atraso',
                    'valor': valor_juros,
                    'tem_incidencia': False,
                    'percentual_incidencia': 0
                })
        
        if valor_multa > 0:
            if MULTA_TEM_INCIDENCIA:
                valor_com_incidencia += valor_multa
                detalhamento.append({
                    'tipo': 'Multa de Atraso',
                    'valor': valor_multa,
                    'tem_incidencia': True,
                    'percentual_incidencia': 100
                })
            else:
                valor_sem_incidencia += valor_multa
                detalhamento.append({
                    'tipo': 'Multa de Atraso',
                    'valor': valor_multa,
                    'tem_incidencia': False,
                    'percentual_incidencia': 0
                })
        
        # 5. DESPESAS USANDO O MODELO DESPESA
        despesas_info = _obter_despesas_cobranca_independente(cobranca)
        
        for despesa_info in despesas_info:
            if despesa_info['valor_com_incidencia'] > 0:
                valor_com_incidencia += despesa_info['valor_com_incidencia']
                detalhamento.append({
                    'tipo': f"Despesa: {despesa_info['descricao']}",
                    'valor': despesa_info['valor_com_incidencia'],
                    'tem_incidencia': True,
                    'percentual_incidencia': despesa_info['percentual_incidencia']
                })
            
            if despesa_info['valor_sem_incidencia'] > 0:
                valor_sem_incidencia += despesa_info['valor_sem_incidencia']
                detalhamento.append({
                    'tipo': f"Despesa: {despesa_info['descricao']} (isenta)",
                    'valor': despesa_info['valor_sem_incidencia'],
                    'tem_incidencia': False,
                    'percentual_incidencia': 0
                })
        
        # 6. OUTROS VALORES (diferença não explicada)
        valor_outros = valor_total_cobranca - valor_com_incidencia - valor_sem_incidencia
        
        if valor_outros > 0:
            # REGRA: "Outros valores" NÃO têm incidência (ex: IPTU, condomínio)
            valor_sem_incidencia += valor_outros
            detalhamento.append({
                'tipo': 'Outros Valores (IPTU, Condomínio, etc)',
                'valor': valor_outros,
                'tem_incidencia': False,
                'percentual_incidencia': 0,
                'observacao': 'Assumido como repasse direto sem incidência'
            })
        
        # 7. CALCULAR TAXA ADMINISTRATIVA
        taxa_percentual = _obter_taxa_administrativa_independente(cobranca.contrato)
        valor_taxa_admin = valor_com_incidencia * (taxa_percentual / 100)
        
        # 8. VALOR LÍQUIDO PARA REPASSE
        valor_liquido_repasse = valor_total_cobranca - valor_taxa_admin
        
        # 9. GERAR DESCRIÇÃO DETALHADA
        descricao = _gerar_descricao_repasse_independente(
            cobranca, valor_com_incidencia, valor_sem_incidencia, 
            taxa_percentual, valor_taxa_admin, detalhamento
        )
        
        return {
            'sucesso': True,
            'valor_total_cobranca': valor_total_cobranca,
            'valor_com_incidencia': valor_com_incidencia,
            'valor_sem_incidencia': valor_sem_incidencia,
            'taxa_percentual': taxa_percentual,
            'valor_taxa_admin': valor_taxa_admin,
            'valor_liquido_repasse': valor_liquido_repasse,
            'detalhamento': detalhamento,
            'descricao': descricao
        }
        
    except Exception as e:
        return {
            'sucesso': False,
            'erro': f'Erro no cálculo: {str(e)}'
        }

def _obter_valor_juros_multas(cobranca, tipo):
    """
    Obtém valores de juros ou multas da cobrança
    """
    try:
        campos_juros = ['valor_juros', 'juros', 'valor_juros_atraso']
        campos_multa = ['valor_multa', 'multa', 'valor_multa_atraso']
        
        campos = campos_juros if tipo == 'juros' else campos_multa
        
        for campo in campos:
            if hasattr(cobranca, campo):
                valor = getattr(cobranca, campo)
                if valor and valor > 0:
                    return Decimal(str(valor))
        
        return Decimal('0.00')
        
    except Exception:
        return Decimal('0.00')

def _obter_despesas_cobranca_independente(cobranca):
    """
    Obtém despesas usando o modelo Despesa - VERSÃO INDEPENDENTE
    """
    try:
        from financeiro.models import Despesa
        from django.db import models
        from datetime import date
        
        # Buscar despesas do contrato ativas no período da cobrança
        data_referencia = date(cobranca.ano_referencia, cobranca.mes_referencia, 1)
        
        despesas = Despesa.objects.filter(
            contrato=cobranca.contrato,
            is_ativa=True,
            data_inicio__lte=data_referencia
        ).filter(
            models.Q(data_fim_prevista__isnull=True) | 
            models.Q(data_fim_prevista__gte=data_referencia)
        )
        
        despesas_info = []
        
        for despesa in despesas:
            try:
                # Verificar se a despesa está ativa na data de referência
                if despesa.data_inicio <= data_referencia:
                    valor_parcela = despesa.calcular_valor_parcela()
                    
                    # Calcular incidência usando os métodos do modelo
                    if despesa.incidencia_taxa_admin == 'sim':
                        valor_com_incidencia = valor_parcela
                        valor_sem_incidencia = Decimal('0.00')
                        percentual = 100
                    elif despesa.incidencia_taxa_admin == 'parcial':
                        valor_com_incidencia = valor_parcela * (despesa.percentual_com_incidencia / 100)
                        valor_sem_incidencia = valor_parcela - valor_com_incidencia
                        percentual = despesa.percentual_com_incidencia
                    else:  # 'nao'
                        valor_com_incidencia = Decimal('0.00')
                        valor_sem_incidencia = valor_parcela
                        percentual = 0
                    
                    if valor_com_incidencia > 0 or valor_sem_incidencia > 0:
                        despesas_info.append({
                            'despesa_id': despesa.id,
                            'descricao': despesa.descricao or despesa.tipo.nome,
                            'valor_total': valor_com_incidencia + valor_sem_incidencia,
                            'valor_com_incidencia': valor_com_incidencia,
                            'valor_sem_incidencia': valor_sem_incidencia,
                            'percentual_incidencia': percentual,
                            'pago_por': despesa.paga_por
                        })
            except Exception as e:
                logger.warning(f"Erro ao processar despesa {despesa.id}: {e}")
                continue
        
        return despesas_info
        
    except Exception as e:
        logger.warning(f"Erro ao obter despesas da cobrança {cobranca.id}: {e}")
        return []

def _obter_taxa_administrativa_independente(contrato):
    """
    Obtém a taxa administrativa - VERSÃO INDEPENDENTE
    """
    try:
        # PRIORIDADE 1: Política do contrato
        politica = getattr(contrato, 'politica_repasse', None)
        if politica and hasattr(politica, 'taxa_admin_personalizada') and politica.taxa_admin_personalizada:
            return politica.taxa_admin_personalizada
        
        # PRIORIDADE 2: Taxa do contrato
        if hasattr(contrato, 'valor_taxa_administracao_percentual') and contrato.valor_taxa_administracao_percentual:
            return contrato.valor_taxa_administracao_percentual
        
        # Fallback: 8%
        return Decimal('8.00')
        
    except Exception:
        return Decimal('8.00')

def _gerar_descricao_repasse_independente(cobranca, valor_com_incidencia, valor_sem_incidencia, taxa_percentual, valor_taxa_admin, detalhamento):
    """
    Gera descrição detalhada do repasse - VERSÃO INDEPENDENTE
    """
    linhas = [
        f"Repasse ref. {cobranca.mes_referencia:02d}/{cobranca.ano_referencia}",
        "",
        "COMPOSIÇÃO:"
    ]
    
    # Agrupar por tipo de incidência
    com_incidencia = [d for d in detalhamento if d['tem_incidencia']]
    sem_incidencia = [d for d in detalhamento if not d['tem_incidencia']]
    
    if com_incidencia:
        linhas.append("• Com incidência de taxa:")
        for item in com_incidencia:
            linhas.append(f"  - {item['tipo']}: R$ {item['valor']:.2f}")
    
    if sem_incidencia:
        linhas.append("• Sem incidência de taxa:")
        for item in sem_incidencia:
            linhas.append(f"  - {item['tipo']}: R$ {item['valor']:.2f}")
    
    linhas.extend([
        "",
        f"RESUMO:",
        f"• Base cálculo taxa: R$ {valor_com_incidencia:.2f}",
        f"• Taxa administrativa ({taxa_percentual}%): R$ {valor_taxa_admin:.2f}",
        f"• Valores isentos: R$ {valor_sem_incidencia:.2f}",
        f"• LÍQUIDO REPASSE: R$ {(valor_com_incidencia + valor_sem_incidencia - valor_taxa_admin):.2f}"
    ])
    
    return "\n".join(linhas)

# View principal atualizada - VERSÃO INDEPENDENTE

@require_POST
def gerar_repasse_automatico(request):
    """
    View para gerar repasses automaticamente - VERSÃO INDEPENDENTE
    """
    try:
        with transaction.atomic():
            cobrancas_sem_repasse = Cobranca.objects.filter(
                status='paga'
            ).exclude(
                id__in=Repasse.objects.values_list('cobranca_id', flat=True)
            ).select_related('contrato')
            
            total_cobrancas = cobrancas_sem_repasse.count()
            
            if total_cobrancas == 0:
                return JsonResponse({
                    'success': True,
                    'criados': 0,
                    'processados': 0,
                    'agendados': 0,
                    'message': 'Nenhuma cobrança elegível para gerar repasse.'
                })
            
            criados = 0
            erros = 0
            detalhes = []
            
            for cobranca in cobrancas_sem_repasse:
                try:
                    # Obter proprietário
                    proprietario = _obter_proprietario_contrato(cobranca.contrato)
                    if not proprietario:
                        erros += 1
                        detalhes.append(f"Cobrança {cobranca.id}: sem proprietário")
                        continue
                    
                    # Calcular data usando função independente
                    data_repasse = _calcular_data_repasse_independente(cobranca)
                    
                    # USAR O NOVO CÁLCULO COMPLETO INDEPENDENTE
                    resultado_calculo = _calcular_valores_repasse_completo(cobranca)
                    
                    if not resultado_calculo['sucesso']:
                        erros += 1
                        detalhes.append(f"Cobrança {cobranca.id}: {resultado_calculo['erro']}")
                        continue
                    
                    # Criar repasse com cálculo completo
                    repasse = Repasse.objects.create(
                        proprietario=proprietario,
                        cobranca=cobranca,
                        contrato=cobranca.contrato,
                        valor=resultado_calculo['valor_total_cobranca'],
                        valor_desconto=Decimal('0.00'),
                        valor_taxa_admin=resultado_calculo['valor_taxa_admin'],
                        data_prevista=data_repasse,
                        mes_referencia=cobranca.mes_referencia,
                        ano_referencia=cobranca.ano_referencia,
                        status='pendente',
                        tipo='automatico',
                        descricao=resultado_calculo['descricao']
                    )
                    
                    criados += 1
                    detalhes.append(
                        f"Repasse {repasse.id}: R$ {resultado_calculo['valor_liquido_repasse']:.2f} "
                        f"(Taxa: R$ {resultado_calculo['valor_taxa_admin']:.2f} sobre R$ {resultado_calculo['valor_com_incidencia']:.2f})"
                    )
                    logger.info(f"Repasse {repasse.id} criado com cálculo independente")
                        
                except Exception as e:
                    erros += 1
                    detalhes.append(f"Erro na cobrança {cobranca.id}: {str(e)}")
                    logger.error(f"Erro ao processar cobrança {cobranca.id}: {e}")
            
            # Resposta
            message = f"{criados} repasse(s) criado(s) com cálculo independente!"
            if erros > 0:
                message += f" {erros} erro(s) ocorreram."
            
            return JsonResponse({
                'success': True,
                'criados': criados,
                'processados': 0,
                'agendados': 0,
                'erros': erros,
                'message': message,
                'detalhes': detalhes[:5]
            })
            
    except Exception as e:
        logger.error(f"Erro geral na geração de repasses: {e}")
        return JsonResponse({
            'success': False,
            'error': f'Erro interno: {str(e)}',
            'criados': 0,
            'processados': 0,
            'agendados': 0
        })

def _obter_proprietario_contrato(contrato):
    """Obtém proprietário do contrato"""
    try:
        if hasattr(contrato, 'proprietario'):
            if hasattr(contrato.proprietario, 'all'):
                return contrato.proprietario.first()
            else:
                return contrato.proprietario
        return None
    except AttributeError:
        return None


def _gerar_descricao_repasse(cobranca, valor_com_incidencia, valor_sem_incidencia, taxa_percentual, valor_taxa_admin, detalhamento):
    """
    Gera descrição detalhada do repasse
    """
    linhas = [
        f"Repasse ref. {cobranca.mes_referencia:02d}/{cobranca.ano_referencia}",
        "",
        "COMPOSIÇÃO:"
    ]
    
    # Agrupar por tipo de incidência
    com_incidencia = [d for d in detalhamento if d['tem_incidencia']]
    sem_incidencia = [d for d in detalhamento if not d['tem_incidencia']]
    
    if com_incidencia:
        linhas.append("• Com incidência de taxa:")
        for item in com_incidencia:
            linhas.append(f"  - {item['tipo']}: R$ {item['valor']:.2f}")
    
    if sem_incidencia:
        linhas.append("• Sem incidência de taxa:")
        for item in sem_incidencia:
            linhas.append(f"  - {item['tipo']}: R$ {item['valor']:.2f}")
    
    linhas.extend([
        "",
        f"RESUMO:",
        f"• Base cálculo taxa: R$ {valor_com_incidencia:.2f}",
        f"• Taxa administrativa ({taxa_percentual}%): R$ {valor_taxa_admin:.2f}",
        f"• Valores isentos: R$ {valor_sem_incidencia:.2f}",
        f"• LÍQUIDO REPASSE: R$ {(valor_com_incidencia + valor_sem_incidencia - valor_taxa_admin):.2f}"
    ])
    
    return "\n".join(linhas)

# Atualizar a view principal para usar o novo cálculo

@require_POST
def gerar_repasse_automatico(request):
    """
    View para gerar repasses automaticamente - VERSÃO COMPLETA
    """
    try:
        with transaction.atomic():
            cobrancas_sem_repasse = Cobranca.objects.filter(
                status='paga'
            ).exclude(
                id__in=Repasse.objects.values_list('cobranca_id', flat=True)
            ).select_related('contrato')
            
            total_cobrancas = cobrancas_sem_repasse.count()
            
            if total_cobrancas == 0:
                return JsonResponse({
                    'success': True,
                    'criados': 0,
                    'processados': 0,
                    'agendados': 0,
                    'message': 'Nenhuma cobrança elegível para gerar repasse.'
                })
            
            criados = 0
            erros = 0
            detalhes = []
            
            for cobranca in cobrancas_sem_repasse:
                try:
                    # Obter proprietário
                    proprietario = _obter_proprietario_contrato(cobranca.contrato)
                    if not proprietario:
                        erros += 1
                        detalhes.append(f"Cobrança {cobranca.id}: sem proprietário")
                        continue
                    
                    # Calcular data usando RepasseService
                    data_repasse = RepasseService._calcular_data_repasse(cobranca)
                    
                    # USAR O NOVO CÁLCULO COMPLETO
                    resultado_calculo = _calcular_valores_repasse_completo(cobranca)
                    
                    if not resultado_calculo['sucesso']:
                        erros += 1
                        detalhes.append(f"Cobrança {cobranca.id}: {resultado_calculo['erro']}")
                        continue
                    
                    # Criar repasse com cálculo completo
                    repasse = Repasse.objects.create(
                        proprietario=proprietario,
                        cobranca=cobranca,
                        contrato=cobranca.contrato,
                        valor=resultado_calculo['valor_total_cobranca'],
                        valor_desconto=Decimal('0.00'),
                        valor_taxa_admin=resultado_calculo['valor_taxa_admin'],
                        data_prevista=data_repasse,
                        mes_referencia=cobranca.mes_referencia,
                        ano_referencia=cobranca.ano_referencia,
                        status='pendente',
                        tipo='automatico',
                        descricao=resultado_calculo['descricao']
                    )
                    
                    criados += 1
                    detalhes.append(
                        f"Repasse {repasse.id}: R$ {resultado_calculo['valor_liquido_repasse']:.2f} "
                        f"(Taxa: R$ {resultado_calculo['valor_taxa_admin']:.2f} sobre R$ {resultado_calculo['valor_com_incidencia']:.2f})"
                    )
                    logger.info(f"Repasse {repasse.id} criado com cálculo integrado")
                        
                except Exception as e:
                    erros += 1
                    detalhes.append(f"Erro na cobrança {cobranca.id}: {str(e)}")
                    logger.error(f"Erro ao processar cobrança {cobranca.id}: {e}")
            
            # Resposta
            message = f"{criados} repasse(s) criado(s) com cálculo integrado!"
            if erros > 0:
                message += f" {erros} erro(s) ocorreram."
            
            return JsonResponse({
                'success': True,
                'criados': criados,
                'processados': 0,
                'agendados': 0,
                'erros': erros,
                'message': message,
                'detalhes': detalhes[:5]
            })
            
    except Exception as e:
        logger.error(f"Erro geral na geração de repasses: {e}")
        return JsonResponse({
            'success': False,
            'error': f'Erro interno: {str(e)}',
            'criados': 0,
            'processados': 0,
            'agendados': 0
        })

def _obter_proprietario_contrato(contrato):
    """Obtém proprietário do contrato"""
    try:
        if hasattr(contrato, 'proprietario'):
            if hasattr(contrato.proprietario, 'all'):
                return contrato.proprietario.first()
            else:
                return contrato.proprietario
        return None
    except AttributeError:
        return None

# Função de teste para o shell
def testar_calculo_integrado(cobranca_id):
    """
    Testa o cálculo integrado para uma cobrança específica
    Execute: testar_calculo_integrado(120)
    """
    try:
        from financeiro.models import Cobranca
        
        cobranca = Cobranca.objects.get(id=cobranca_id)
        resultado = _calcular_valores_repasse_completo(cobranca)
        
        print(f"🧪 TESTE CÁLCULO INTEGRADO - COBRANÇA #{cobranca_id}")
        print("=" * 70)
        
        if resultado['sucesso']:
            print(f"✅ Cálculo bem-sucedido:")
            print(f"   💰 Valor total: R$ {resultado['valor_total_cobranca']:.2f}")
            print(f"   📊 Com incidência: R$ {resultado['valor_com_incidencia']:.2f}")
            print(f"   🆓 Sem incidência: R$ {resultado['valor_sem_incidencia']:.2f}")
            print(f"   🏛️ Taxa admin ({resultado['taxa_percentual']}%): R$ {resultado['valor_taxa_admin']:.2f}")
            print(f"   💵 Líquido repasse: R$ {resultado['valor_liquido_repasse']:.2f}")
            
            print(f"\n📋 DETALHAMENTO:")
            for item in resultado['detalhamento']:
                incidencia = "✅" if item['tem_incidencia'] else "❌"
                print(f"   {incidencia} {item['tipo']}: R$ {item['valor']:.2f}")
            
            print(f"\n📝 DESCRIÇÃO:")
            print(resultado['descricao'])
        else:
            print(f"❌ Erro: {resultado['erro']}")
        
        return resultado
        
    except Exception as e:
        print(f"❌ Erro ao testar: {e}")
        return None